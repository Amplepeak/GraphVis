"""Folder-scale batch processing and report generation for GraphVis."""
from __future__ import annotations

import html
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

import pandas as pd

from graphvis.core.logging import get_logger
from graphvis.data.loader import load_dataset

LOG = get_logger("batch")


@dataclass(slots=True)
class BatchItemResult:
    path: str
    ok: bool
    summary: str
    payload: Any = None


@dataclass(slots=True)
class BatchRunResult:
    items: list[BatchItemResult] = field(default_factory=list)
    output_files: list[str] = field(default_factory=list)

    @property
    def successes(self) -> int: return sum(i.ok for i in self.items)
    @property
    def failures(self) -> int: return len(self.items)-self.successes


class BatchProcessor:
    DATA_EXTENSIONS={".csv",".tsv",".txt",".asc",".dat",".json",".xlsx",".xls",".xlsm",".mat",".h5",".hdf",".hdf5",".nc",".netcdf",".tdms",".tdm",".html",".htm",".xml",".wav",".mzml",".mzxml",".mgf",".fcs"}

    def scan(self,folder: str,recursive: bool = False) -> list[str]:
        root=Path(folder); it=root.rglob("*") if recursive else root.glob("*"); return [str(p) for p in it if p.is_file() and p.suffix.lower() in self.DATA_EXTENSIONS]

    def run(self,paths: Iterable[str],operation: Callable[[Any],Any] | None = None) -> BatchRunResult:
        out=BatchRunResult()
        for path in paths:
            try:
                ds=load_dataset(path); payload=operation(ds) if operation else {"rows":len(ds.df),"columns":list(ds.df.columns),"description":ds.describe()}; out.items.append(BatchItemResult(path,True,ds.describe(),payload))
            except Exception as exc:
                LOG.exception("Batch item failed: %s",path); out.items.append(BatchItemResult(path,False,f"{type(exc).__name__}: {exc}"))
        return out

    @staticmethod
    def write_html_report(result: BatchRunResult,path: str,title: str = "GraphVis Batch Report") -> str:
        rows=[]
        for item in result.items:
            payload=html.escape(json.dumps(item.payload,default=str,ensure_ascii=False)[:4000]) if item.ok else ""
            rows.append(f"<tr><td>{html.escape(Path(item.path).name)}</td><td>{'OK' if item.ok else 'ERROR'}</td><td>{html.escape(item.summary)}</td><td><pre>{payload}</pre></td></tr>")
        doc=f"""<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>body{{font-family:Arial,sans-serif;margin:2em}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #ccc;padding:.5em;vertical-align:top}}pre{{white-space:pre-wrap}}</style></head><body><h1>{html.escape(title)}</h1><p>{result.successes} succeeded; {result.failures} failed.</p><table><tr><th>File</th><th>Status</th><th>Summary</th><th>Payload</th></tr>{''.join(rows)}</table></body></html>"""
        Path(path).write_text(doc,encoding="utf-8"); result.output_files.append(path); return path

    @staticmethod
    def write_word_report(result: BatchRunResult,path: str,title: str = "GraphVis Batch Report") -> str:
        try:
            from docx import Document
        except Exception as exc: raise RuntimeError("Word report generation requires python-docx.") from exc
        doc=Document(); doc.add_heading(title,0); doc.add_paragraph(f"{result.successes} succeeded; {result.failures} failed."); table=doc.add_table(rows=1,cols=3); hdr=table.rows[0].cells; hdr[0].text="File"; hdr[1].text="Status"; hdr[2].text="Summary"
        for item in result.items:
            cells=table.add_row().cells; cells[0].text=Path(item.path).name; cells[1].text="OK" if item.ok else "ERROR"; cells[2].text=item.summary
        doc.save(path); result.output_files.append(path); return path

    @staticmethod
    def write_pdf_report(result: BatchRunResult,path: str,title: str = "GraphVis Batch Report") -> str:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet
        except Exception as exc: raise RuntimeError("PDF report generation requires reportlab.") from exc
        styles=getSampleStyleSheet(); story=[Paragraph(title,styles["Title"]),Spacer(1,12),Paragraph(f"{result.successes} succeeded; {result.failures} failed.",styles["BodyText"]),Spacer(1,12)]
        data=[["File","Status","Summary"]]+[[Path(i.path).name,"OK" if i.ok else "ERROR",i.summary[:180]] for i in result.items]
        table=Table(data,repeatRows=1,colWidths=[140,55,300]); table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.25,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey),("VALIGN",(0,0),(-1,-1),"TOP")]))
        story.append(table); SimpleDocTemplate(path,pagesize=A4).build(story); result.output_files.append(path); return path
