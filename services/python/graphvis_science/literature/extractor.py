# =========================================================================
# literature.py — end-to-end literature extraction.
#
#   extract_literature(path)  ->  LiteratureExtraction
#
# Pipeline (each stage is optional and degrades gracefully):
#   1. text + tables via pdfplumber (lines + text strategies)
#   2. text + tables via PyMuPDF  (page.find_tables)
#   3. OCR via PyMuPDF rasterisation + pytesseract when no text layer exists
#   4. numeric series recovered from the running text
#        - whitespace-aligned numeric blocks (with header sniffing)
#        - "value ± error" pairs (kept as separate value / err columns)
#        - loose x, y number pairs (last resort)
#   5. operational parameters, units, ranges and experimental conditions
#   6. datasets saved to literature_dataset/{title}_dataset_{n}.csv, with a
#      {title}_extraction.json side-car describing provenance + parameters
#
# Pure Python; safe to run in a worker thread.
# =========================================================================
from __future__ import annotations

import json
import os
import re
import time
from dataclasses import dataclass, field

import numpy as np
import pandas as pd

from graphvis_science.runtime import LITERATURE_DIR, get_logger

LOG = get_logger("literature_extractor")
# The path is known at import time; the folder is only created when something
# actually writes to it - `save_extraction` does its own makedirs. Nothing in
# runtime does a mkdir, because a redirected or offline Documents folder used
# to make importing this module raise and report the whole add-on broken.
LITERATURE_DATASET_DIR = str(LITERATURE_DIR)

_NUM = r"[-+]?(?:\d+\.\d*|\.\d+|\d+)(?:[eE][-+]?\d+)?"
_NUM_RE = re.compile(_NUM)
_PM = r"(?:±|\+/-|\+-|\u00b1)"


@dataclass
class LitDataset:
    name: str
    df: pd.DataFrame
    source_page: int | None = None
    kind: str = "table"          # table | text_block | pm_pairs | number_pairs
    kind_hint: str = "generic"   # gompertz | polarisation | nyquist | bode | time_series | generic
    header_note: str = ""
    yerr_col: str | None = None
    saved_path: str | None = None


@dataclass
class LiteratureExtraction:
    title: str
    path: str
    method: str = ""
    text: str = ""
    datasets: list[LitDataset] = field(default_factory=list)
    parameters: dict = field(default_factory=dict)
    saved_paths: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    n_pages: int = 0
    primary_document: str | None = None
    linked_files: list[str] = field(default_factory=list)
    batch_id: str = ""
    semantic_context: dict = field(default_factory=dict)

    def summary(self) -> str:
        lines = [f"Title: {self.title}", f"Method: {self.method or 'n/a'}; pages: {self.n_pages}; text chars: {len(self.text):,}",
                 f"Datasets extracted: {len(self.datasets)}"]
        for d in self.datasets:
            lines.append(f"  • {d.name} — {len(d.df)}×{len(d.df.columns)} [{d.kind}, hint: {d.kind_hint}]"
                         + (f" p.{d.source_page}" if d.source_page else ""))
        if self.parameters:
            lines.append("Parameters: " + "; ".join(f"{k} = {v['value']} {v.get('unit', '')}".strip()
                                                    for k, v in self.parameters.items()))
        if self.linked_files:
            lines.append(f"Linked files: {len(self.linked_files)}" + (f"; primary: {os.path.basename(self.primary_document)}" if self.primary_document else ""))
        if self.semantic_context:
            sc = self.semantic_context
            lines.append(f"Semantic context: {sc.get('inferred_plot_type','n/a')} (confidence {float(sc.get('confidence',0)):.2f}; {sc.get('backend','heuristic')})")
            axes = [f"X={sc.get('x_variable')}" if sc.get('x_variable') else "", f"Y={sc.get('y_variable')}" if sc.get('y_variable') else "", f"Z={sc.get('z_variable')}" if sc.get('z_variable') else ""]
            axes = [a for a in axes if a]
            if axes: lines.append("Semantic axes: " + ", ".join(axes))
            if sc.get('conditions'): lines.append("Context conditions: " + "; ".join(f"{k}={v.get('value')} {v.get('unit','')}".strip() for k,v in list(sc['conditions'].items())[:8]))
        if self.warnings:
            lines.append("Warnings: " + " | ".join(self.warnings))
        return "\n".join(lines)


# ------------------------------------------------------------------ title
def sanitise_title(title: str, max_len: int = 60) -> str:
    t = re.sub(r"[^A-Za-z0-9 _\-]+", " ", str(title or "")).strip()
    t = re.sub(r"\s+", "_", t)
    return (t[:max_len].rstrip("_") or "literature")


def _title_from_text(text: str, fallback: str) -> str:
    for line in text.splitlines()[:25]:
        s = line.strip()
        if 15 <= len(s) <= 180 and not _NUM_RE.fullmatch(s) and s.count(" ") >= 2:
            if not re.match(r"^(abstract|keywords|received|accepted|available online|journal|©|doi)", s, re.I):
                return s
    return fallback


# ------------------------------------------------------------ pdf readers
def _read_pdf_pdfplumber(path, progress, *, deadline=None, cancelled=None):
    import pdfplumber
    text, tables, n_pages = [], [], 0
    with pdfplumber.open(path) as pdf:
        n_pages = len(pdf.pages)
        for i, page in enumerate(pdf.pages, 1):
            if cancelled and cancelled():
                raise RuntimeError("Cancelled")
            if deadline is not None and time.monotonic() >= deadline:
                break
            if progress:
                progress(f"pdfplumber: page {i}/{n_pages}", int(10 + 40 * i / max(n_pages, 1)))
            pt = page.extract_text() or ""
            text.append(pt)
            seen = set()
            for settings in ({}, {"vertical_strategy": "text", "horizontal_strategy": "text"}):
                try:
                    found = page.extract_tables(table_settings=settings) if settings else page.extract_tables()
                except Exception:
                    found = []
                for tb in found or []:
                    if not tb or len(tb) < 2:
                        continue
                    sig = (len(tb), len(tb[0]), str(tb[0])[:60])
                    if sig in seen:
                        continue
                    seen.add(sig)
                    tables.append((i, tb))
    meta_title = ""
    try:
        with pdfplumber.open(path) as pdf:
            meta_title = (pdf.metadata or {}).get("Title", "") or ""
    except Exception:
        LOG.debug("no PDF metadata title for %s", path, exc_info=True)
    return "\n".join(text), tables, n_pages, meta_title


def _read_pdf_pymupdf(path, progress, *, deadline=None, cancelled=None):
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore
    text, tables = [], []
    doc = fitz.open(path)
    n_pages = doc.page_count
    for i, page in enumerate(doc, 1):
        if cancelled and cancelled():
            doc.close(); raise RuntimeError("Cancelled")
        if deadline is not None and time.monotonic() >= deadline:
            break
        if progress:
            progress(f"PyMuPDF: page {i}/{n_pages}", int(50 + 20 * i / max(n_pages, 1)))
        text.append(page.get_text() or "")
        if hasattr(page, "find_tables"):
            try:
                for tb in page.find_tables().tables:
                    rows = tb.extract()
                    if rows and len(rows) >= 2:
                        tables.append((i, rows))
            except Exception:
                LOG.debug("table detection failed on a page of %s", path, exc_info=True)
    meta_title = (doc.metadata or {}).get("title", "") or ""
    doc.close()
    return "\n".join(text), tables, n_pages, meta_title


def _ocr_pdf(path, progress, dpi=200, max_pages=None, *, deadline=None, cancelled=None):
    """Rasterise pages with PyMuPDF and OCR them with pytesseract."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return "", "pytesseract/Pillow not installed"
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz  # type: ignore
        except ImportError:
            return "", "PyMuPDF not installed (needed to rasterise pages for OCR)"
    doc = fitz.open(path)
    out = []
    n = doc.page_count if max_pages is None else min(doc.page_count, max_pages)
    for i in range(n):
        if cancelled and cancelled():
            doc.close(); raise RuntimeError("Cancelled")
        if deadline is not None and time.monotonic() >= deadline:
            break
        if progress:
            progress(f"OCR: page {i + 1}/{n}", int(50 + 40 * (i + 1) / max(n, 1)))
        pix = doc[i].get_pixmap(dpi=dpi)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        try:
            out.append(pytesseract.image_to_string(img))
        except Exception as exc:
            doc.close()
            return "\n".join(out), f"OCR failed: {exc}"
    doc.close()
    return "\n".join(out), ""


# ------------------------------------------------------- table utilities
def _to_number(cell):
    if cell is None:
        return np.nan
    s = str(cell).strip().replace(",", "")
    s = s.replace("−", "-").replace("–", "-")
    m = _NUM_RE.search(s)
    if not m:
        return np.nan
    try:
        return float(m.group(0))
    except ValueError:
        return np.nan


def _split_pm(cell):
    if cell is None:
        return np.nan, np.nan
    s = str(cell).replace("−", "-")
    m = re.match(rf"\s*({_NUM})\s*{_PM}\s*({_NUM})", s)
    if m:
        return float(m.group(1)), float(m.group(2))
    return _to_number(s), np.nan


def _clean_header(h, idx):
    s = re.sub(r"\s+", " ", str(h or "")).strip().replace("\n", " ")
    return s if s else f"col_{idx}"


def table_to_dataframe(rows: list[list]) -> pd.DataFrame | None:
    """Convert a raw table (list of rows) into a numeric DataFrame.
    Header = first row; "x ± e" cells become x and x_err columns; columns
    with < 50 % numeric cells are dropped."""
    if not rows or len(rows) < 2:
        return None
    width = max(len(r) for r in rows)
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    header = [_clean_header(h, i) for i, h in enumerate(rows[0])]
    body = rows[1:]
    data = {}
    for j, name in enumerate(header):
        col = [r[j] for r in body]
        has_pm = any(re.search(_PM, str(c or "")) for c in col)
        if has_pm:
            vals, errs = zip(*[_split_pm(c) for c in col])
            data[name] = list(vals)
            data[f"{name}_err"] = list(errs)
        else:
            data[name] = [_to_number(c) for c in col]
    df = pd.DataFrame(data)
    df = df.dropna(how="all")
    keep = [c for c in df.columns if df[c].notna().mean() >= 0.5]
    df = df[keep]
    # de-duplicate column names
    seen, cols = {}, []
    for c in df.columns:
        if c in seen:
            seen[c] += 1
            cols.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 0
            cols.append(c)
    df.columns = cols
    if len(df) < 2 or len(df.columns) < 2:
        return None
    return df.reset_index(drop=True)


# ------------------------------------------------- numeric blocks in text
_CELL_RE = re.compile(rf"({_NUM})(?:\s*{_PM}\s*({_NUM}))?")


def _numeric_cells(line):
    """Return [(value, err|nan), ...] if the line is purely numeric cells
    (optionally 'value ± err'), else None."""
    s = line.replace("−", "-").replace("–", "-").strip()
    if not s:
        return None
    cells, residue = [], _CELL_RE.sub(" ", s)
    if re.sub(r"[\s(),;%\[\]]", "", residue):
        return None
    for m in _CELL_RE.finditer(s):
        try:
            cells.append((float(m.group(1)), float(m.group(2)) if m.group(2) else np.nan))
        except ValueError:
            return None
    return cells if len(cells) >= 2 else None


def numeric_blocks_from_text(text: str, min_rows=3) -> list[tuple[pd.DataFrame, str]]:
    """Find runs of ≥ min_rows consecutive lines that each contain the same
    number (≥2) of numeric cells ('12.5' or '12.5 ± 1.1'). The preceding
    line is used as a header when its field count matches."""
    lines = text.splitlines()
    out = []
    i = 0
    while i < len(lines):
        cells = _numeric_cells(lines[i])
        if cells:
            width = len(cells)
            j, block = i, []
            while j < len(lines):
                cj = _numeric_cells(lines[j])
                if cj and len(cj) == width:
                    block.append(cj)
                    j += 1
                else:
                    break
            if len(block) >= min_rows:
                header = None
                if i > 0:
                    raw = lines[i - 1].strip()
                    hw = [w for w in re.split(r"\s{2,}|\t", raw) if w]
                    if len(hw) != width:
                        hw = [w for w in re.split(r"\)\s+", raw) if w]
                        hw = [w if w.endswith(")") or "(" not in w else w + ")" for w in hw]
                    if len(hw) == width:
                        header = [_clean_header(w, k) for k, w in enumerate(hw)]
                header = header or [f"col_{k}" for k in range(width)]
                data = {}
                for k, name in enumerate(header):
                    data[name] = [row[k][0] for row in block]
                    errs = [row[k][1] for row in block]
                    if np.isfinite(errs).any():
                        data[f"{name}_err"] = errs
                df = pd.DataFrame(data)
                out.append((df, lines[i - 1].strip() if i > 0 else ""))
            i = max(j, i + 1)
        else:
            i += 1
    return out


def _values_signature(df: pd.DataFrame) -> set:
    v = df.select_dtypes(include=[np.number]).to_numpy(float).ravel()
    return {round(float(x), 6) for x in v if np.isfinite(x)}


def _header_quality(df: pd.DataFrame) -> int:
    return sum(1 for c in df.columns if re.search(r"[A-Za-z]{3,}", str(c)) and not str(c).startswith("col_"))


def dedupe_datasets(datasets: list) -> list:
    """Drop datasets whose numeric content is (almost) contained in another
    one, keeping the richer / better-labelled version."""
    keep = []
    sigs = [(_values_signature(d.df), d) for d in datasets]
    for i, (si, di) in enumerate(sigs):
        dominated = False
        for j, (sj, dj) in enumerate(sigs):
            if i == j or not si:
                continue
            overlap = len(si & sj) / len(si)
            if overlap >= 0.8:
                better = (len(sj) > len(si)) or (len(sj) == len(si) and (_header_quality(dj.df), len(dj.df.columns), -j) > (_header_quality(di.df), len(di.df.columns), -i))
                if better:
                    dominated = True
                    break
        if not dominated:
            keep.append(di)
    return keep


def pm_pairs_from_text(text: str, min_rows=3) -> pd.DataFrame | None:
    """Recover 'x  y ± e' style rows scattered through the text."""
    rows = re.findall(rf"({_NUM})[\s,;]+({_NUM})\s*{_PM}\s*({_NUM})", text.replace("−", "-"))
    if len(rows) < min_rows:
        return None
    arr = np.array(rows, dtype=float)
    return pd.DataFrame(arr, columns=["X_extracted", "Y_extracted", "Y_extracted_err"])


def number_pairs_from_text(text: str, min_rows=6) -> pd.DataFrame | None:
    pairs = re.findall(rf"(?<![\w.])({_NUM})[,\s]+({_NUM})(?![\w.])", text.replace("−", "-"))
    if len(pairs) < min_rows:
        return None
    arr = np.array(pairs, dtype=float)
    # drop obviously non-data pairs (years / page refs)
    ok = ~((arr[:, 0] >= 1900) & (arr[:, 0] <= 2100) & (np.mod(arr[:, 0], 1) == 0))
    arr = arr[ok]
    if len(arr) < min_rows:
        return None
    return pd.DataFrame(arr, columns=["X_extracted", "Y_extracted"])


# ---------------------------------------------------------- classification
def classify_columns(columns) -> str:
    cols = " ".join(str(c).lower() for c in columns)
    if ("z'" in cols or "zre" in cols or "z_re" in cols or "real" in cols) and ("z''" in cols or "zim" in cols or "z_im" in cols or "imag" in cols):
        return "nyquist"
    if "freq" in cols and ("phase" in cols or "|z|" in cols or "zmod" in cols):
        return "bode"
    if ("current" in cols or "j (" in cols or "a/m" in cols or "ma/cm" in cols) and ("volt" in cols or "e (" in cols or "potential" in cols):
        return "polarisation"
    if ("h2" in cols or "hydrogen" in cols or "cumulative" in cols) and ("time" in cols or "day" in cols or "hour" in cols or " h" in cols):
        return "gompertz"
    if "time" in cols or "day" in cols or "hour" in cols or "hrt" in cols:
        return "time_series"
    return "generic"


# ----------------------------------------------------------- parameters
_PARAM_PATTERNS = {
    "voltage": (rf"(?:applied\s+(?:voltage|potential)|cell\s+voltage|potential)[^0-9\-]{{0,40}}({_NUM})\s*(m?V)\b", "V"),
    "pH": (rf"\bpH\s*(?:of|=|:|was|at)?\s*({_NUM})", ""),
    "temperature": (rf"({_NUM})\s*(?:°|º|deg)\s*C\b", "°C"),
    "retention_time": (rf"(?:retention\s+time|residence\s+time)[^0-9]{{0,30}}({_NUM})\s*(s|sec|seconds?|min|minutes?|h|hr|hours?|d|days?)\b", ""),
    "current_density": (rf"current\s+density[^0-9]{{0,30}}({_NUM})\s*(A\s*/?\s*m[²2]|mA\s*/?\s*cm[²2]|A\s*m-2)", ""),
    "electrode_area": (rf"(?:electrode|surface)\s+area[^0-9]{{0,30}}({_NUM})\s*(cm[²2]|m[²2])", ""),
    "pressure": (rf"(?:pressure)[^0-9\-]{{0,30}}({_NUM})\s*(Pa|kPa|MPa|bar|mbar|atm)\b", ""),
    "frequency": (rf"(?:frequency)[^0-9\-]{{0,30}}({_NUM})\s*(Hz|kHz|MHz)\b", ""),
}


def extract_parameters(text: str) -> dict:
    out = {}
    flat = re.sub(r"\s+", " ", text)
    for key, (pat, default_unit) in _PARAM_PATTERNS.items():
        m = re.search(pat, flat, re.I)
        if not m:
            continue
        try:
            value = float(m.group(1))
        except (ValueError, IndexError):
            continue
        unit = (m.group(2) if m.lastindex and m.lastindex >= 2 else default_unit) or default_unit
        start = max(m.start() - 60, 0)
        out[key] = {"value": value, "unit": re.sub(r"\s+", "", unit or ""),
                    "context": flat[start:m.end() + 40].strip()}
    return out


# ------------------------------------------------------------- storage
def _next_dataset_number(title: str, output_dir: str | None = None) -> int:
    n = 0
    pat = re.compile(rf"^{re.escape(title)}_dataset_(\d+)\.csv$")
    target = output_dir or LITERATURE_DATASET_DIR
    if os.path.isdir(target):
        for f in os.listdir(target):
            m = pat.match(f)
            if m:
                n = max(n, int(m.group(1)))
    return n + 1


def save_extraction(ext: LiteratureExtraction, output_dir: str | None = None) -> list[str]:
    target = os.path.abspath(output_dir or LITERATURE_DATASET_DIR)
    os.makedirs(target, exist_ok=True)
    title = sanitise_title(ext.title)
    n = _next_dataset_number(title, target)
    saved = []
    for d in ext.datasets:
        path = os.path.join(target, f"{title}_dataset_{n}.csv")
        d.df.to_csv(path, index=False)
        d.saved_path = path
        saved.append(path)
        n += 1
    side = {
        "title": ext.title, "source": ext.path, "method": ext.method, "pages": ext.n_pages,
        "primary_document": ext.primary_document, "linked_files": ext.linked_files, "batch_id": ext.batch_id,
        "parameters": ext.parameters,
        "semantic_context": ext.semantic_context,
        "datasets": [{"file": os.path.basename(d.saved_path), "kind": d.kind, "hint": d.kind_hint,
                      "page": d.source_page, "columns": list(map(str, d.df.columns)), "rows": int(len(d.df)),
                      "header_note": d.header_note} for d in ext.datasets],
        "warnings": ext.warnings,
    }
    try:
        with open(os.path.join(target, f"{title}_extraction.json"), "w", encoding="utf-8") as fh:
            json.dump(side, fh, indent=2, ensure_ascii=False)
    except Exception:
        LOG.debug("could not write the extraction summary for %s", title, exc_info=True)
    ext.saved_paths = saved
    return saved


def plottability(df) -> dict:
    """How much of this table can actually be drawn.

    WHY THIS EXISTS
    ---------------
    "Send extracted data to workspace" took the FIRST of the extracted
    datasets and imported it. A paper yielded twelve, and the first was a
    reference list: pdfplumber had read a block of prose as a table, so the
    columns came out as clipped word fragments ("nvironme", "y and E") and the
    one column that looked numeric was a run of publication years. The figure
    drew a single point, and the person reasonably reported that the button did
    not work.

    It had worked exactly as written. "First" is not a choice, it is the
    absence of one, and in a PDF the first thing that parses as a table is very
    often the thing that is not data.

    What makes a table drawable is not its kind, its name or its position: it
    is whether at least two of its columns are numbers, on enough rows to make
    a line. That is measured here rather than guessed, with the SAME coercion
    the plotting path uses, so a table this calls drawable is one the canvas
    can draw.

    A column counts as numeric when at least four fifths of its non-empty cells
    parse as finite numbers - not all of them, because a real table of results
    carries the occasional "n/a", "<0.01" or footnote marker and throwing the
    whole column away for one of those would reject the data and keep the
    reference list.
    """
    import pandas as pd

    numeric_names, rows_per_column = [], {}
    for name in df.columns:
        # dropna FIRST, then stringify.
        #
        # The other way round turns a genuinely missing cell into the string
        # "None" or "nan", which is then counted as a present non-numeric cell -
        # so a column of numbers with a few gaps in it was scored as prose. An
        # own test caught this: two float columns with real gaps reported zero
        # numeric columns between them.
        #
        # A cell that is EMPTY is absent. A cell that says "n/a" or carries a
        # footnote marker is present and not a number, which is what the four
        # fifths rule is there to tolerate.
        filled = df[name].dropna().astype(str).str.strip()
        filled = filled[filled != ""]
        if len(filled) == 0:
            continue
        coerced = pd.to_numeric(filled, errors="coerce")
        good = int(coerced.notna().sum())
        if good >= 0.8 * len(filled) and good >= 2:
            numeric_names.append(str(name))
            rows_per_column[str(name)] = good

    # Rows where at least TWO numeric columns are both present, which is what a
    # point on a figure needs. Counted on the frame rather than summed per
    # column, so a table whose two numeric columns never overlap scores zero.
    plottable = 0
    if len(numeric_names) >= 2:
        block = df[numeric_names].apply(pd.to_numeric, errors="coerce")
        plottable = int((block.notna().sum(axis=1) >= 2).sum())

    # HOW MANY COLUMNS HAVE A NAME.
    #
    # Reported alongside the numbers because the numbers alone chose a table of
    # contents. Twelve tables out of a 69-page thesis arrived with headers like
    # "demic", "i", "of Fi", "List of E", "4.1.5", "col_13" - pdfplumber
    # reading a contents page, a running head or a fragment of body text, with
    # no header row to find. col_N is its own placeholder for exactly that.
    #
    # The decision is made in AppController, on the saved file, because the
    # add-on is installed as a copy and may be older than the app. This is the
    # same measurement so the two agree about what they are looking at.
    named = 0
    for name in df.columns:
        text = str(name).strip()
        if len(text) < 3 or text.startswith("col_"):
            continue
        try:
            float(text)
        except ValueError:
            named += 1
    return {"numeric_columns": len(numeric_names),
            "numeric_names": numeric_names,
            "named_columns": named,
            "plottable_rows": plottable}


def load_extraction_sidecars() -> list[dict]:
    out = []
    if not os.path.isdir(LITERATURE_DATASET_DIR):
        return out
    for f in sorted(os.listdir(LITERATURE_DATASET_DIR)):
        if f.endswith("_extraction.json"):
            try:
                with open(os.path.join(LITERATURE_DATASET_DIR, f), encoding="utf-8") as fh:
                    out.append(json.load(fh))
            except Exception:
                LOG.debug("skipping unreadable extracted dataset %s", f, exc_info=True)
    return out


# --------------------------------------------------------------- driver
def _strip_markup(raw: str) -> str:
    """Plain text out of XML or HTML, without adding a parser dependency.

    Script and style bodies go first - their contents are not prose and would
    otherwise dominate the word counts the semantic analysis works from.
    """
    import html as _html
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<(br|/p|/div|/tr|/h[1-6])[^>]*>", "\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    raw = _html.unescape(raw)
    raw = re.sub(r"[ \t]+", " ", raw)
    return re.sub(r"\n\s*\n\s*\n+", "\n\n", raw).strip()

def extract_literature(path: str, extract_dataset: bool = True, ocr: bool = True,
                       progress=None, output_dir: str | None = None,
                       time_budget_seconds: float | None = None, cancelled=None) -> LiteratureExtraction:
    def report(msg, pct=-1):
        if progress:
            progress(msg, pct)

    stem = os.path.splitext(os.path.basename(path))[0]
    ext = LiteratureExtraction(title=stem, path=path)
    budget = None if time_budget_seconds is None else max(10.0, min(float(time_budget_seconds), 259200.0))
    deadline = None if budget is None else time.monotonic() + budget

    def check_cancelled() -> None:
        if cancelled and cancelled():
            raise RuntimeError("Cancelled")

    def budget_expired() -> bool:
        return deadline is not None and time.monotonic() >= deadline

    def note_budget() -> None:
        if budget_expired() and not any("time budget" in str(w).lower() for w in ext.warnings):
            ext.warnings.append("Literature time budget reached; the extracted content up to this save point was retained and can be reused later.")

    lower = path.lower()
    text, tables, meta_title, n_pages = "", [], "", 0
    methods = []

    if lower.endswith(".pdf"):
        report("Opening PDF…", 5)
        check_cancelled()
        try:
            text, tables, n_pages, meta_title = _read_pdf_pdfplumber(path, report, deadline=deadline, cancelled=cancelled)
            methods.append("pdfplumber")
        except ImportError:
            ext.warnings.append("pdfplumber not installed")
        except RuntimeError as exc:
            if str(exc) == "Cancelled":
                raise
            ext.warnings.append(f"pdfplumber failed: {exc}")
        except Exception as exc:
            ext.warnings.append(f"pdfplumber failed: {exc}")
        try:
            check_cancelled()
            if budget_expired():
                raise TimeoutError("literature time budget reached")
            t2, tables2, n2, mt2 = _read_pdf_pymupdf(path, report, deadline=deadline, cancelled=cancelled)
            methods.append("pymupdf")
            n_pages = n_pages or n2
            meta_title = meta_title or mt2
            if len(t2) > len(text) * 1.2:
                text = t2
            sigs = {(p, len(tb), len(tb[0])) for p, tb in tables}
            for p, tb in tables2:
                if (p, len(tb), len(tb[0])) not in sigs:
                    tables.append((p, tb))
        except ImportError:
            if not methods:
                ext.warnings.append("PyMuPDF not installed")
        except TimeoutError:
            note_budget()
        except RuntimeError as exc:
            if str(exc) == "Cancelled":
                raise
            ext.warnings.append(f"PyMuPDF failed: {exc}")
        except Exception as exc:
            ext.warnings.append(f"PyMuPDF failed: {exc}")
        if ocr and not budget_expired() and len(re.sub(r"\s+", "", text)) < 200 * max(n_pages, 1) * 0.05:
            report("No text layer found — running OCR…", 50)
            ocr_text, err = _ocr_pdf(path, report, deadline=deadline, cancelled=cancelled)
            if ocr_text.strip():
                text = ocr_text if len(ocr_text) > len(text) else text
                methods.append("ocr")
            if err:
                ext.warnings.append(err)
        note_budget()
    elif lower.endswith((".csv", ".tsv")):
        try:
            df = pd.read_csv(path, sep=None, engine="python")
            for c in df.columns:
                df[c] = pd.to_numeric(df[c], errors="coerce")
            df = df.dropna(axis=1, how="all")
            if len(df.columns) >= 2:
                tables.append((None, [list(df.columns)] + df.values.tolist()))
            methods.append("csv")
        except Exception as exc:
            ext.warnings.append(f"CSV read failed: {exc}")
    elif lower.endswith((".xlsx", ".xls")):
        try:
            df = pd.read_excel(path)
            for c in df.columns:
                df[c] = pd.to_numeric(df[c], errors="coerce")
            df = df.dropna(axis=1, how="all")
            if len(df.columns) >= 2:
                tables.append((None, [list(df.columns)] + df.values.tolist()))
            methods.append("excel")
        except Exception as exc:
            ext.warnings.append(f"Excel read failed: {exc}")
    elif lower.endswith(".json"):
        try:
            try:
                df = pd.read_json(path)
            except ValueError:
                df = pd.read_json(path, lines=True)
            for c in df.columns:
                df[c] = pd.to_numeric(df[c], errors="coerce")
            df = df.dropna(axis=1, how="all")
            if len(df.columns) >= 2:
                tables.append((None, [list(df.columns)] + df.values.tolist()))
            methods.append("json")
        except Exception as exc:
            ext.warnings.append(f"JSON read failed: {exc}")
    elif lower.endswith((".docx",)):
        # A thesis chapter or a lab protocol is as often a Word document as a
        # PDF, and refusing it made "read the literature" narrower than the work
        # it is meant to support. Tables come across as tables, not as prose.
        try:
            from docx import Document as _Docx
            doc = _Docx(path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            for table in doc.tables:
                rows = [[c.text for c in r.cells] for r in table.rows]
                if len(rows) > 1 and len(rows[0]) >= 2:
                    tables.append((None, rows))
            methods.append("docx")
        except Exception as exc:
            ext.warnings.append(f"Word read failed: {exc}")
    elif lower.endswith((".odt", ".epub")):
        # Both are zip archives of XML. Unpacking and stripping the markup needs
        # no extra dependency, and the alternative packages disagree about
        # whitespace anyway.
        try:
            import zipfile
            parts = []
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                wanted = ([n for n in names if n.endswith("content.xml")]
                          if lower.endswith(".odt")
                          else [n for n in names if n.lower().endswith((".xhtml", ".html", ".htm"))])
                for name in wanted:
                    parts.append(archive.read(name).decode("utf-8", errors="ignore"))
            text = _strip_markup("\n".join(parts))
            methods.append("odt" if lower.endswith(".odt") else "epub")
        except Exception as exc:
            ext.warnings.append(f"Archive read failed: {exc}")
    elif lower.endswith((".html", ".htm", ".xhtml")):
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                raw = fh.read()
            text = _strip_markup(raw)
            # A page of results usually has its numbers in a real table.
            try:
                for df in pd.read_html(path):
                    numeric = df.apply(pd.to_numeric, errors="coerce").dropna(axis=1, how="all")
                    if len(numeric.columns) >= 2:
                        tables.append((None, [list(df.columns)] + df.values.tolist()))
            except Exception:
                LOG.debug("read_html found nothing usable in %s", path, exc_info=True)
            methods.append("html")
        except Exception as exc:
            ext.warnings.append(f"HTML read failed: {exc}")
    elif lower.endswith(".rtf"):
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                raw = fh.read()
            # RTF is control words, groups and escaped characters. This keeps the
            # prose and drops the formatting; it is not a full parser and does
            # not pretend to be.
            raw = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: chr(int(m.group(1), 16)), raw)
            raw = re.sub(r"\\par[d]?\b", "\n", raw)
            raw = re.sub(r"\\[a-zA-Z]+-?\d*\s?", " ", raw)
            text = re.sub(r"[{}]", " ", raw)
            text = re.sub(r"[ \t]+", " ", text).strip()
            methods.append("rtf")
        except Exception as exc:
            ext.warnings.append(f"RTF read failed: {exc}")
    else:
        check_cancelled()
        with open(path, "r", encoding="utf-8", errors="ignore") as fh:
            text = fh.read()
        methods.append("text")

    ext.text = text
    ext.n_pages = n_pages
    ext.method = "+".join(methods)
    ext.title = (meta_title.strip() if meta_title and len(meta_title.strip()) > 5 else _title_from_text(text, stem))

    check_cancelled()
    report("Parsing tables…", 75)
    idx = 0
    for page, rows in tables:
        if budget_expired():
            note_budget(); break
        check_cancelled()
        df = table_to_dataframe(rows)
        if df is None:
            continue
        idx += 1
        ext.datasets.append(LitDataset(name=f"table_{idx}", df=df, source_page=page, kind="table",
                                       kind_hint=classify_columns(df.columns),
                                       header_note=" | ".join(map(str, rows[0]))[:120]))

    report("Scanning text for numeric series…", 85)
    if text and not budget_expired():
        check_cancelled()
        for df, note in numeric_blocks_from_text(text):
            if len(df) >= 3:
                idx += 1
                ext.datasets.append(LitDataset(name=f"text_block_{idx}", df=df, kind="text_block",
                                               kind_hint=classify_columns(df.columns), header_note=note[:120]))
        if not ext.datasets:
            pm = pm_pairs_from_text(text)
            if pm is not None:
                idx += 1
                ext.datasets.append(LitDataset(name=f"pm_pairs_{idx}", df=pm, kind="pm_pairs",
                                               kind_hint="generic", yerr_col="Y_extracted_err"))
            else:
                pairs = number_pairs_from_text(text)
                if pairs is not None:
                    idx += 1
                    ext.datasets.append(LitDataset(name=f"number_pairs_{idx}", df=pairs, kind="number_pairs"))
        ext.parameters = extract_parameters(text)

    ext.datasets = dedupe_datasets(ext.datasets)
    # Context-aware NLP pass: captions, abstract, experimental conditions,
    # variable relationships and likely graph/axis assignments. The heuristic
    # core works offline; a local Transformers model can optionally refine it.
    try:
        check_cancelled()
        if budget_expired():
            raise TimeoutError("literature time budget reached")
        from graphvis_science.literature.semantics import analyse_literature_context
        context = analyse_literature_context(text, [d.df for d in ext.datasets])
        ext.semantic_context = context.to_dict()
        semantic_hint_map = {
            "Modified Gompertz Kinetics": "gompertz", "Polarisation & Power Curve": "polarisation",
            "EIS: Nyquist": "nyquist", "EIS: Bode": "bode", "Line Chart": "time_series",
            "2D Heatmap": "heatmap", "2D Contour": "contour", "3D Topography / Surface": "surface",
            "Histogram": "histogram", "Box Plot": "box", "Violin Plot": "violin",
            "Pareto Front": "pareto", "4D / 5D Scatter": "scatter",
        }
        semantic_hint = semantic_hint_map.get(context.inferred_plot_type)
        if semantic_hint:
            for d in ext.datasets:
                if d.kind_hint in ("generic", "time_series") or context.confidence >= 0.70:
                    d.kind_hint = semantic_hint
    except TimeoutError:
        note_budget()
    except Exception as exc:
        ext.warnings.append(f"Semantic context pass failed: {type(exc).__name__}: {exc}")

    for n, d in enumerate(ext.datasets, 1):
        d.name = re.sub(r"_\d+$", f"_{n}", d.name)
        if d.yerr_col is None:
            errs = [c for c in d.df.columns if str(c).endswith("_err")]
            d.yerr_col = errs[0] if errs else None

    if extract_dataset and ext.datasets:
        check_cancelled()
        report("Saving datasets…", 95)
        save_extraction(ext, output_dir=output_dir)
    note_budget()
    report("Done", 100)
    return ext
