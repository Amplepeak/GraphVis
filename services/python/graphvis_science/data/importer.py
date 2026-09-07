"""Comprehensive dataset import.

The native Rust core reads Arrow IPC, CSV/TSV and Parquet directly. Everything
else is converted here into Arrow IPC and handed back to it - the "optional
scientific IO plugin that emits Arrow IPC" the core's own error message
describes.

Readers are registered by extension in READERS below, so adding a format means
adding one function. Every third-party reader is imported lazily: a format
whose package is not installed reports exactly what to install instead of
failing at start-up, and never blocks the formats that do work.

Ported from GraphVis 17 data/loader.py and extended well beyond it.
"""
from __future__ import annotations

import os
import shlex
import shutil
import sqlite3
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd


class ImportError_(Exception):
    """Raised with a message intended for the user."""


def _need(module: str, package: str):
    try:
        __import__(module)
    except Exception as exc:  # pragma: no cover - depends on the install
        raise ImportError_(
            f"This format needs the optional '{package}' package "
            f"(import of {module} failed: {exc})."
        ) from exc


def _frame_from_arrays(arrays: dict[str, np.ndarray]) -> pd.DataFrame:
    """Align 1-D numeric arrays into one frame.

    GraphVis 17 keeps arrays matching the modal length and truncates longer
    ones, so a file mixing a long time base with shorter derived series still
    yields a usable table.
    """
    series = {}
    for k, v in arrays.items():
        a = np.asarray(v)
        if a.ndim == 1 and a.size >= 2:
            series[str(k)] = a
    if not series:
        return pd.DataFrame()
    lengths = [len(v) for v in series.values()]
    modal = max(set(lengths), key=lengths.count)
    kept = {k: v[:modal] for k, v in series.items() if len(v) >= modal}
    if not kept:
        return pd.DataFrame()
    frame = pd.DataFrame(kept)
    # Drop only columns that carry no information at all. This used to drop
    # every column with a single distinct value, which quietly deleted setpoint
    # channels, flat baselines, run ids and flags from MATLAB, HDF5, TDMS, EDF,
    # ABF, NWB, seismic and xarray imports - the column simply was not in the
    # Variable Mapping list and nothing said why.
    keep = [c for c in frame.columns if frame[c].notna().any()]
    return frame[keep] if keep else frame


def _largest(frames) -> pd.DataFrame:
    frames = [f for f in frames if f is not None and len(f)]
    if not frames:
        raise ImportError_("No tabular data was found in this file.")
    return max(frames, key=len)


# ---------------------------------------------------------------- plain tables
def _r_csv(p):     return pd.read_csv(p)
def _r_tsv(p):     return pd.read_csv(p, sep='\t')
def _r_sniff(p):
    # Instrument text dumps use anything from commas to runs of spaces.
    try:
        return pd.read_csv(p, sep=None, engine='python')
    except Exception:
        return pd.read_csv(p, sep=r'\s+', engine='python')
def _r_json(p):    return pd.read_json(p)
def _r_jsonl(p):   return pd.read_json(p, lines=True)
def _r_excel(p):
    return _largest(pd.read_excel(p, sheet_name=None).values())
def _r_xlsb(p):
    _need('pyxlsb', 'pyxlsb')
    return _largest(pd.read_excel(p, sheet_name=None, engine='pyxlsb').values())
def _r_ods(p):
    _need('odf', 'odfpy')
    return _largest(pd.read_excel(p, sheet_name=None, engine='odf').values())
def _r_html(p):
    _need('lxml', 'lxml')
    return _largest(pd.read_html(p))
def _r_xml(p):
    _need('lxml', 'lxml')
    return pd.read_xml(p)
def _r_yaml(p):
    _need('yaml', 'pyyaml')
    import yaml
    with open(p, 'r', encoding='utf-8') as fh:
        return pd.json_normalize(yaml.safe_load(fh))
def _r_toml(p):
    try:
        import tomllib
    except Exception:
        _need('tomli', 'tomli'); import tomli as tomllib
    with open(p, 'rb') as fh:
        return pd.json_normalize(tomllib.load(fh))

# ------------------------------------------------------------------- columnar
def _r_parquet(p):
    _need('pyarrow', 'pyarrow'); return pd.read_parquet(p)
def _r_feather(p):
    _need('pyarrow', 'pyarrow'); return pd.read_feather(p)
def _r_orc(p):
    _need('pyarrow', 'pyarrow'); return pd.read_orc(p)
def _r_arrow(p):
    import pyarrow.ipc as ipc, pyarrow as pa
    with pa.memory_map(str(p), 'r') as src:
        return ipc.open_file(src).read_all().to_pandas()

# -------------------------------------------------------- statistics packages
def _r_spss(p):    _need('pyreadstat', 'pyreadstat'); return pd.read_spss(p)
def _r_stata(p):   return pd.read_stata(p)
def _r_sas(p):     return pd.read_sas(p)

# ---------------------------------------------------------------- numpy/pickle
def _r_npy(p):
    a = np.load(p, allow_pickle=False)
    a = np.squeeze(a)
    if a.ndim == 1:
        return pd.DataFrame({'value': a})
    if a.ndim == 2:
        return pd.DataFrame(a, columns=[f'col_{i}' for i in range(a.shape[1])])
    raise ImportError_(f"{a.ndim}-D arrays cannot become a table; export 1-D or 2-D.")
def _r_npz(p):
    with np.load(p, allow_pickle=False) as z:
        return _frame_from_arrays({k: z[k] for k in z.files})
def _r_pickle(p):
    obj = pd.read_pickle(p)
    if isinstance(obj, pd.DataFrame):
        return obj
    if isinstance(obj, dict):
        return _frame_from_arrays({k: np.asarray(v) for k, v in obj.items()
                                   if isinstance(v, (list, tuple, np.ndarray))})
    raise ImportError_("This pickle does not contain a DataFrame or arrays.")

# --------------------------------------------------------------------- sqlite
def _r_sqlite(p):
    con = sqlite3.connect(str(p))
    try:
        names = pd.read_sql_query(
            "SELECT name FROM sqlite_master WHERE type IN ('table','view')", con)['name'].tolist()
        if not names:
            raise ImportError_("This database has no tables.")
        return _largest([pd.read_sql_query(f'SELECT * FROM "{n}"', con) for n in names])
    finally:
        con.close()

# ------------------------------------------------------------------ MATLAB/HDF
def _hdf5_series(path) -> dict[str, np.ndarray]:
    """Every 1-D numeric dataset in an HDF5 file, keyed by its leaf name.

    MATLAB v7.3 files are HDF5, so the .mat reader and the .h5/.hdf5 reader
    walked the file with the same nine lines. One walk now: a change to what
    counts as a usable series - the dtype kinds, the minimum length - lands in
    both formats instead of one.
    """
    _need('h5py', 'h5py')
    import h5py
    out: dict[str, np.ndarray] = {}

    def visit(name, node):
        if isinstance(node, h5py.Dataset) and node.dtype.kind in ('f', 'i', 'u', 'b'):
            arr = np.squeeze(np.asarray(node))
            if arr.ndim == 1 and arr.size >= 2:
                out[name.split('/')[-1]] = arr.astype(np.float64, copy=False)

    with h5py.File(path, 'r') as fh:
        fh.visititems(visit)
    return out


def _r_matlab(p):
    _need('h5py', 'h5py')
    import h5py
    raw: dict[str, np.ndarray] = {}
    if h5py.is_hdf5(p):                       # MATLAB v7.3 is HDF5
        raw = _hdf5_series(p)
    else:                                     # v5/v6/v7
        _need('scipy', 'scipy')
        from scipy.io import loadmat
        data = loadmat(p, squeeze_me=True, struct_as_record=False)

        def walk(prefix, obj):
            if isinstance(obj, np.ndarray) and obj.dtype.kind in ('f', 'i', 'u', 'b'):
                arr = np.squeeze(obj)
                if arr.ndim == 1 and arr.size >= 2:
                    raw[prefix.split('/')[-1]] = arr.astype(np.float64, copy=False)
            elif hasattr(obj, '_fieldnames'):
                for fn in obj._fieldnames:
                    walk(f"{prefix}/{fn}", getattr(obj, fn))
            elif isinstance(obj, np.ndarray) and obj.dtype == object:
                for i, item in enumerate(obj.ravel()):
                    walk(f"{prefix}/{i}", item)
        for k, v in data.items():
            if not str(k).startswith('__'):
                walk(str(k), v)
    return _frame_from_arrays(raw)

def _r_hdf(p):
    return _frame_from_arrays(_hdf5_series(p))

def _r_netcdf(p):
    _need('xarray', 'xarray')
    import xarray as xr
    ds = xr.open_dataset(p)
    try:
        return ds.to_dataframe().reset_index()
    finally:
        ds.close()

# ------------------------------------------------------------ instruments/DAQ
def _r_tdms(p):
    _need('nptdms', 'nptdms')
    from nptdms import TdmsFile
    cols: dict[str, np.ndarray] = {}
    with TdmsFile.open(p) as tdms:
        for group in tdms.groups():
            for ch in group.channels():
                try:
                    data = np.asarray(ch[:])
                except Exception:
                    continue
                if data.ndim == 1 and data.size >= 2:
                    cols[f"{group.name}/{ch.name}" if group.name else ch.name] = data
    return _frame_from_arrays(cols)

def _r_tdm(p):
    side = Path(p).with_suffix('.tdx')
    if not side.exists() and not Path(p).with_suffix('.TDX').exists():
        raise ImportError_("A .tdm file needs its matching .tdx data file beside it.")
    raise ImportError_("TDM/TDX is not supported yet; export as TDMS or CSV.")

def _r_audio(p):
    _need('soundfile', 'soundfile')
    import soundfile as sf
    data, rate = sf.read(p, always_2d=True)
    cols = {'time_s': np.arange(data.shape[0], dtype=np.float64) / float(rate)}
    for c in range(data.shape[1]):
        cols[f'channel_{c + 1}'] = data[:, c].astype(np.float64)
    return pd.DataFrame(cols)

def _r_las(p):
    _need('lasio', 'lasio')
    import lasio
    return lasio.read(p).df().reset_index()

# ------------------------------------------------------------------ mass spec
def _r_massspec(p):
    _need('pyteomics', 'pyteomics')
    low = str(p).lower()
    if low.endswith('.mgf'):
        from pyteomics import mgf as mod
    elif low.endswith('.mzml'):
        from pyteomics import mzml as mod
    else:
        from pyteomics import mzxml as mod
    mz, it, scan = [], [], []
    with mod.read(str(p)) as spectra:
        for i, spec in enumerate(spectra):
            m = np.asarray(spec.get('m/z array', []), dtype=np.float64)
            v = np.asarray(spec.get('intensity array', []), dtype=np.float64)
            if m.size and m.size == v.size:
                mz.append(m); it.append(v); scan.append(np.full(m.size, i, dtype=np.float64))
    if not mz:
        raise ImportError_("No spectra with m/z and intensity arrays were found.")
    return pd.DataFrame({'scan': np.concatenate(scan),
                         'mz': np.concatenate(mz),
                         'intensity': np.concatenate(it)})

def _r_fcs(p):
    _need('fcsparser', 'fcsparser')
    import fcsparser
    _meta, frame = fcsparser.parse(str(p), reformat_meta=True)
    return frame

def _r_jcamp(p):
    _need('jcamp', 'jcamp')
    from jcamp import jcamp_readfile
    d = jcamp_readfile(str(p))
    x, y = np.asarray(d.get('x', [])), np.asarray(d.get('y', []))
    if x.size and x.size == y.size:
        return pd.DataFrame({str(d.get('xunits', 'x')): x, str(d.get('yunits', 'y')): y})
    raise ImportError_("This JCAMP-DX file has no matching x/y arrays.")

# ------------------------------------------------------------------ astronomy
def _r_fits(p):
    _need('astropy', 'astropy')
    from astropy.table import Table
    from astropy.io import fits
    with fits.open(p) as hdus:
        for hdu in hdus:
            if getattr(hdu, 'data', None) is None:
                continue
            try:
                return Table(hdu.data).to_pandas()
            except Exception:
                arr = np.squeeze(np.asarray(hdu.data))
                if arr.ndim == 1:
                    return pd.DataFrame({'value': arr})
                if arr.ndim == 2:
                    return pd.DataFrame(arr, columns=[f'col_{i}' for i in range(arr.shape[1])])
    raise ImportError_("No table or 1-D/2-D image was found in this FITS file.")

def _r_root(p):
    _need('uproot', 'uproot')
    import uproot
    with uproot.open(p) as f:
        for key, obj in f.items():
            try:
                return obj.arrays(library='pd')
            except Exception:
                continue
    raise ImportError_("No readable TTree was found in this ROOT file.")

# ------------------------------------------------------------------- seismic
def _r_seismic(p):
    _need('obspy', 'obspy')
    from obspy import read as obspy_read
    st = obspy_read(str(p))
    cols: dict[str, np.ndarray] = {}
    for i, tr in enumerate(st):
        cols[f"{tr.id or 'trace'}_{i}"] = np.asarray(tr.data, dtype=np.float64)
    if st:
        cols['time_s'] = np.arange(len(st[0].data), dtype=np.float64) * float(st[0].stats.delta)
    return _frame_from_arrays(cols)

# ---------------------------------------------------------- neuro / physiology
def _r_edf(p):
    _need('pyedflib', 'pyedflib')
    import pyedflib
    with pyedflib.EdfReader(str(p)) as fh:
        labels = fh.getSignalLabels()
        cols = {labels[i]: fh.readSignal(i) for i in range(fh.signals_in_file)}
    return _frame_from_arrays(cols)

def _r_mne(p):
    _need('mne', 'mne')
    import mne
    low = str(p).lower()
    if low.endswith('.fif'):
        raw = mne.io.read_raw_fif(p, preload=True, verbose='ERROR')
    elif low.endswith('.set'):
        raw = mne.io.read_raw_eeglab(p, preload=True, verbose='ERROR')
    elif low.endswith('.cnt'):
        raw = mne.io.read_raw_cnt(p, preload=True, verbose='ERROR')
    elif low.endswith('.bdf'):
        raw = mne.io.read_raw_bdf(p, preload=True, verbose='ERROR')
    else:
        raw = mne.io.read_raw(p, preload=True, verbose='ERROR')
    return raw.to_data_frame()

def _r_abf(p):
    _need('pyabf', 'pyabf')
    import pyabf
    abf = pyabf.ABF(str(p))
    cols = {'time_s': np.asarray(abf.sweepX, dtype=np.float64)}
    for ch in range(abf.channelCount):
        abf.setSweep(0, channel=ch)
        cols[abf.adcNames[ch] if ch < len(abf.adcNames) else f'channel_{ch}'] = \
            np.asarray(abf.sweepY, dtype=np.float64)
    return _frame_from_arrays(cols)

def _r_nwb(p):
    _need('pynwb', 'pynwb')
    from pynwb import NWBHDF5IO
    with NWBHDF5IO(str(p), 'r', load_namespaces=True) as io:
        nwb = io.read()
        cols: dict[str, np.ndarray] = {}
        for name, series in getattr(nwb, 'acquisition', {}).items():
            data = np.squeeze(np.asarray(getattr(series, 'data', [])))
            if data.ndim == 1 and data.size >= 2:
                cols[str(name)] = data.astype(np.float64, copy=False)
        return _frame_from_arrays(cols)

# ---------------------------------------------------------------- medical img
def _r_dicom(p):
    _need('pydicom', 'pydicom')
    import pydicom
    ds = pydicom.dcmread(str(p))
    arr = np.squeeze(np.asarray(ds.pixel_array))
    if arr.ndim == 1:
        return pd.DataFrame({'value': arr})
    if arr.ndim == 2:
        return pd.DataFrame(arr, columns=[f'col_{i}' for i in range(arr.shape[1])])
    raise ImportError_("Only 1-D or 2-D DICOM pixel data becomes a table.")

def _r_nifti(p):
    _need('nibabel', 'nibabel')
    import nibabel as nib
    arr = np.squeeze(np.asarray(nib.load(str(p)).dataobj))
    if arr.ndim == 1:
        return pd.DataFrame({'value': arr})
    if arr.ndim == 2:
        return pd.DataFrame(arr, columns=[f'col_{i}' for i in range(arr.shape[1])])
    raise ImportError_("Only 1-D or 2-D NIfTI data becomes a table.")

def _r_tiff(p):
    _need('tifffile', 'tifffile')
    import tifffile
    arr = np.squeeze(tifffile.imread(str(p)))
    if arr.ndim == 2:
        return pd.DataFrame(arr, columns=[f'col_{i}' for i in range(arr.shape[1])])
    raise ImportError_("Only a single 2-D TIFF plane becomes a table.")

# ----------------------------------------------------------------- geospatial
def _r_geo(p):
    _need('geopandas', 'geopandas')
    import geopandas as gpd
    gdf = gpd.read_file(p)
    frame = pd.DataFrame(gdf.drop(columns=gdf.geometry.name, errors='ignore'))
    try:
        pts = gdf.geometry.representative_point()
        frame['longitude'] = pts.x.to_numpy()
        frame['latitude'] = pts.y.to_numpy()
    except Exception:
        pass
    return frame



# ------------------------------------------------------- business and finance
def _r_dbf(p):
    # dBase tables still underpin a lot of accounting, ERP and GIS exports.
    _need('dbfread', 'dbfread')
    from dbfread import DBF
    return pd.DataFrame(iter(DBF(str(p), load=True, char_decode_errors='ignore')))

def _r_access(p):
    """Microsoft Access. Uses the ACE/Jet ODBC driver on Windows."""
    _need('pyodbc', 'pyodbc')
    import pyodbc
    drivers = [d for d in pyodbc.drivers() if 'Access' in d]
    if not drivers:
        raise ImportError_(
            "No Microsoft Access ODBC driver was found. Install the Microsoft "
            "Access Database Engine redistributable, or export the table to CSV."
        )
    con = pyodbc.connect(f"DRIVER={{{drivers[0]}}};DBQ={os.path.abspath(p)};")
    try:
        names = [r.table_name for r in con.cursor().tables(tableType='TABLE')]
        if not names:
            raise ImportError_("This Access database has no user tables.")
        return _largest([pd.read_sql(f'SELECT * FROM [{n}]', con) for n in names])
    finally:
        con.close()

def _r_duckdb(p):
    _need('duckdb', 'duckdb')
    import duckdb
    con = duckdb.connect(str(p), read_only=True)
    try:
        names = [r[0] for r in con.execute("SHOW TABLES").fetchall()]
        if not names:
            raise ImportError_("This DuckDB database has no tables.")
        return _largest([con.execute(f'SELECT * FROM "{n}"').df() for n in names])
    finally:
        con.close()

def _r_avro(p):
    _need('fastavro', 'fastavro')
    import fastavro
    with open(p, 'rb') as fh:
        return pd.DataFrame(list(fastavro.reader(fh)))

def _r_msgpack(p):
    _need('msgpack', 'msgpack')
    import msgpack
    with open(p, 'rb') as fh:
        obj = msgpack.unpack(fh, raw=False)
    if isinstance(obj, list):
        return pd.json_normalize(obj)
    if isinstance(obj, dict):
        return pd.json_normalize([obj])
    raise ImportError_("This MessagePack file does not contain records.")

def _r_rdata(p):
    """R .rds / .RData / .rda."""
    _need('pyreadr', 'pyreadr')
    import pyreadr
    result = pyreadr.read_r(str(p))
    return _largest(list(result.values()))

def _r_numbers(p):
    _need('numbers_parser', 'numbers-parser')
    from numbers_parser import Document
    doc = Document(str(p))
    frames = []
    for sheet in doc.sheets:
        for table in sheet.tables:
            rows = table.rows(values_only=True)
            if len(rows) > 1:
                frames.append(pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]]))
    return _largest(frames)

def _r_fixed_width(p):
    # Column widths are inferred; bank and ledger extracts are often fixed-width.
    return pd.read_fwf(p)

def _r_psv(p):
    return pd.read_csv(p, sep='|')

def _r_ofx(p):
    """OFX / QFX bank and brokerage statements -> one row per transaction."""
    _need('ofxtools', 'ofxtools')
    from ofxtools.Parser import OFXTree
    tree = OFXTree()
    tree.parse(str(p))
    ofx = tree.convert()
    rows = []
    for stmt in list(getattr(ofx, 'statements', []) or []):
        for tx in list(getattr(stmt, 'transactions', []) or []):
            rows.append({
                'date': getattr(tx, 'dtposted', None),
                'amount': float(getattr(tx, 'trnamt', 0) or 0),
                'type': str(getattr(tx, 'trntype', '') or ''),
                'name': str(getattr(tx, 'name', '') or ''),
                'memo': str(getattr(tx, 'memo', '') or ''),
            })
    if not rows:
        raise ImportError_("No transactions were found in this OFX/QFX file.")
    frame = pd.DataFrame(rows)
    frame['date'] = pd.to_datetime(frame['date'], errors='coerce', utc=True)
    return frame

def _r_qif(p):
    """Quicken Interchange Format: records separated by '^'."""
    rows, current = [], {}
    field = {'D': 'date', 'T': 'amount', 'U': 'amount', 'P': 'payee',
             'M': 'memo', 'L': 'category', 'N': 'number'}
    with open(p, 'r', encoding='utf-8', errors='ignore') as fh:
        for line in fh:
            line = line.rstrip('\n').rstrip('\r')
            if not line or line.startswith('!'):
                continue
            if line == '^':
                if current:
                    rows.append(current); current = {}
                continue
            key = field.get(line[0])
            if key:
                current[key] = line[1:].strip()
    if current:
        rows.append(current)
    if not rows:
        raise ImportError_("No transactions were found in this QIF file.")
    frame = pd.DataFrame(rows)
    if 'amount' in frame:
        frame['amount'] = pd.to_numeric(frame['amount'].astype(str).str.replace(',', '', regex=False),
                                        errors='coerce')
    if 'date' in frame:
        frame['date'] = pd.to_datetime(frame['date'], errors='coerce', dayfirst=False)
    return frame

def _r_mt940(p):
    """SWIFT MT940 bank statements."""
    _need('mt940', 'mt-940')
    import mt940 as mt940_mod
    transactions = mt940_mod.parse(str(p))
    rows = []
    for tx in transactions:
        d = dict(tx.data)
        rows.append({
            'date': d.get('date'),
            'entry_date': d.get('entry_date'),
            'amount': float(getattr(d.get('amount'), 'amount', 0) or 0),
            'currency': getattr(d.get('amount'), 'currency', ''),
            'description': str(d.get('transaction_details', '') or ''),
        })
    if not rows:
        raise ImportError_("No transactions were found in this MT940 file.")
    frame = pd.DataFrame(rows)
    frame['date'] = pd.to_datetime(frame['date'], errors='coerce')
    return frame

def _r_pdf_tables(p):
    """Tables inside a PDF - the usual shape of a published financial report."""
    _need('pdfplumber', 'pdfplumber')
    import pdfplumber
    frames = []
    with pdfplumber.open(str(p)) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                if len(table) > 1:
                    header = [str(c) if c is not None else f'col_{i}'
                              for i, c in enumerate(table[0])]
                    frames.append(pd.DataFrame(table[1:], columns=header))
    if not frames:
        raise ImportError_(
            "No tables were detected in this PDF. Use Read Literature for text "
            "and figure extraction instead."
        )
    return _largest(frames)

def _r_docx_tables(p):
    _need('docx', 'python-docx')
    from docx import Document as DocxDocument
    doc = DocxDocument(str(p))
    frames = []
    for table in doc.tables:
        rows = [[c.text for c in r.cells] for r in table.rows]
        if len(rows) > 1:
            frames.append(pd.DataFrame(rows[1:], columns=rows[0]))
    return _largest(frames)

# ------------------------------------------- marine, ocean, diving and survey
def _xr_frame(ds):
    """Flatten an xarray Dataset into a table.

    Ocean and atmosphere files are gridded, so the honest flattening is
    to_dataframe(); it is only refused when the cross product would be too
    large to hold, in which case each variable is reduced to its profile.
    """
    try:
        size = 1
        for length in ds.sizes.values():
            size *= max(int(length), 1)
        if size <= 5_000_000:
            frame = ds.to_dataframe().reset_index()
            if len(frame):
                return frame
    except Exception:
        pass
    arrays = {}
    for name, var in list(ds.coords.items()) + list(ds.data_vars.items()):
        try:
            a = np.asarray(var.values)
            if not np.issubdtype(a.dtype, np.number):
                continue
            while a.ndim > 1:
                a = np.nanmean(a, axis=0)
            if a.ndim == 1:
                arrays[str(name)] = a
        except Exception:
            continue
    return _frame_from_arrays(arrays)


def _read_numeric_text(p, names=None):
    """Read a whitespace/delimiter separated numeric dump that may or may not
    carry a header row."""
    frame = pd.read_csv(p, sep=None, engine='python', header=None, comment='#')
    head = frame.iloc[0]
    if pd.to_numeric(head, errors='coerce').isna().all():
        frame.columns = [str(v).strip() for v in head]
        frame = frame.iloc[1:].reset_index(drop=True)
    elif names:
        frame.columns = (list(names) + [f'col{i}' for i in range(len(frame.columns))])[:len(frame.columns)]
    for col in frame.columns:
        converted = pd.to_numeric(frame[col], errors='coerce')
        if converted.notna().mean() > 0.8:
            frame[col] = converted
    return frame


def _seabird_names(lines):
    names = []
    for raw in lines:
        s = raw.strip()
        if s.lower().startswith('# name ') and '=' in s:
            label = s.split('=', 1)[1].strip()
            # "prDM: Pressure, Digiquartz [db]" -> prDM
            names.append(label.split(':', 1)[0].strip() or f'col{len(names)}')
    return names


def _r_seabird(p):
    """Sea-Bird CTD files (.cnv, .btl, .ctd).

    The instrument writes a '*'/'#' header ending in *END*, in which
    '# name N = short: long [units]' declares each column, followed by
    whitespace separated scan lines. Depth, temperature, salinity and oxygen
    profiles from every CTD cast come through here.
    """
    with open(p, 'r', errors='replace') as fh:
        lines = fh.readlines()
    names = _seabird_names(lines)
    body, started = [], False
    for raw in lines:
        s = raw.strip()
        if not started:
            # A blank line inside the header is not the end of it. Without this
            # the first empty line flipped started=True and every '# name N ='
            # declaration after it was appended to the body as though it were a
            # scan, putting NaN rows through the middle of the cast.
            if not s:
                continue
            if s.startswith('*') or s.startswith('#'):
                if s.upper().startswith('*END*'):
                    started = True
                continue
            started = True
        if s:
            body.append(s.split())
    if not body:
        raise ImportError_("No scan lines were found in this Sea-Bird file.")
    widths = [len(r) for r in body]
    width = max(set(widths), key=widths.count)
    body = [r for r in body if len(r) == width]
    if len(names) != width:
        names = [f'col{i}' for i in range(width)]
    frame = pd.DataFrame(body, columns=names)
    return frame.apply(pd.to_numeric, errors='coerce').dropna(axis=1, how='all')


def _r_odv(p):
    """Ocean Data View spreadsheet export: '//' metadata lines then a tab table.

    ODV is the exchange format for World Ocean Database, GLODAP, SeaDataNet and
    most cruise data centres.
    """
    import io as _io
    with open(p, 'r', errors='replace') as fh:
        lines = fh.readlines()
    start = next((i for i, l in enumerate(lines) if not l.startswith('//')), 0)
    return pd.read_csv(_io.StringIO(''.join(lines[start:])), sep='\t')


def _r_rsk(p):
    """RBR logger archives (.rsk) are SQLite databases; read the largest table."""
    return _r_sqlite(p)


def _r_adcp(p):
    """Acoustic current profilers and velocimeters.

    Teledyne RDI (.pd0, .000) and Nortek (.vec, .aqd, .wpr, .prf) binaries,
    read through dolfyn, which returns velocity, echo intensity and the
    ancillary sensor channels as an xarray Dataset.
    """
    _need('dolfyn', 'dolfyn')
    import dolfyn
    return _xr_frame(dolfyn.read(str(p)))


def _r_grib(p):
    """GRIB/GRIB2 - operational ocean and atmosphere model output (Copernicus
    Marine, HYCOM, ECMWF, NOAA WaveWatch)."""
    _need('cfgrib', 'cfgrib')
    import xarray as xr
    return _xr_frame(xr.open_dataset(str(p), engine='cfgrib'))


def _nmea_degrees(value, hemisphere):
    try:
        raw = float(value)
    except (TypeError, ValueError):
        return np.nan
    degrees = int(raw // 100)
    decimal = degrees + (raw - degrees * 100) / 60.0
    return -decimal if str(hemisphere).upper() in ('S', 'W') else decimal


def _r_nmea(p):
    """NMEA 0183 sentence logs from a survey or dive boat.

    GGA gives the fix, DBT/DPT the echo-sounder depth, MTW water temperature
    and VTG speed over ground - the channels a track plot needs.
    """
    rows = []
    with open(p, 'r', errors='replace') as fh:
        for line in fh:
            s = line.strip()
            if not s.startswith('$') and not s.startswith('!'):
                continue
            f = s.split('*')[0].split(',')
            kind = f[0][-3:].upper()
            if kind == 'GGA' and len(f) > 9:
                row = {'sentence': 'GGA', 'utc': f[1],
                       'latitude': _nmea_degrees(f[2], f[3]),
                       'longitude': _nmea_degrees(f[4], f[5]),
                       'fix_quality': f[6], 'satellites': f[7],
                       'hdop': f[8], 'altitude': f[9]}
            elif kind == 'DBT' and len(f) > 5:
                row = {'sentence': 'DBT', 'depth_feet': f[1],
                       'depth_metres': f[3], 'depth_fathoms': f[5]}
            elif kind == 'DPT' and len(f) > 2:
                row = {'sentence': 'DPT', 'depth_metres': f[1],
                       'transducer_offset': f[2]}
            elif kind == 'MTW' and len(f) > 1:
                row = {'sentence': 'MTW', 'water_temperature': f[1]}
            elif kind == 'VTG' and len(f) > 7:
                row = {'sentence': 'VTG', 'course_true': f[1],
                       'speed_knots': f[5], 'speed_kmh': f[7]}
            else:
                continue
            rows.append(row)
    if not rows:
        raise ImportError_("No usable NMEA sentences were found.")
    return pd.DataFrame(rows)


def _r_gpx(p):
    """GPX tracks - vessel survey lines, shore dives and surface swims."""
    import xml.etree.ElementTree as ET
    root = ET.parse(str(p)).getroot()
    rows = []
    for el in root.iter():
        tag = el.tag.split('}')[-1]
        if tag not in ('trkpt', 'rtept', 'wpt'):
            continue
        row = {'kind': tag}
        try:
            row['latitude'] = float(el.get('lat'))
            row['longitude'] = float(el.get('lon'))
        except (TypeError, ValueError):
            continue
        for child in el.iter():
            name = child.tag.split('}')[-1]
            text = (child.text or '').strip()
            if name != tag and text:
                row.setdefault(name, text)
        rows.append(row)
    if not rows:
        raise ImportError_("This GPX file contains no track, route or waypoint fixes.")
    return pd.DataFrame(rows)


def _r_kml(p):
    """KML/KMZ - dive site marks, transect lines and survey polygons."""
    import zipfile
    import xml.etree.ElementTree as ET
    path = str(p)
    if path.lower().endswith('.kmz'):
        with zipfile.ZipFile(path) as zf:
            inner = next((n for n in zf.namelist() if n.lower().endswith('.kml')), None)
            if inner is None:
                raise ImportError_("This KMZ archive contains no KML document.")
            data = zf.read(inner)
    else:
        with open(path, 'rb') as fh:
            data = fh.read()
    root = ET.fromstring(data)
    rows, feature = [], ''
    for el in root.iter():
        tag = el.tag.split('}')[-1]
        if tag == 'name' and (el.text or '').strip():
            feature = el.text.strip()
        elif tag == 'coordinates' and (el.text or '').strip():
            for triple in el.text.split():
                parts = triple.split(',')
                if len(parts) >= 2:
                    rows.append({'feature': feature,
                                 'longitude': float(parts[0]),
                                 'latitude': float(parts[1]),
                                 'elevation': float(parts[2]) if len(parts) > 2 else np.nan})
    if not rows:
        raise ImportError_("No coordinates were found in this KML document.")
    return pd.DataFrame(rows)


def _r_xyz(p):
    """Two unrelated formats share .xyz, so decide on content.

    A molecular XYZ opens with an integer atom count; a bathymetric or survey
    XYZ is columns of easting/northing/depth soundings.
    """
    with open(p, 'r', errors='replace') as fh:
        first = fh.readline().strip()
    if first.isdigit():
        return _r_xyz_molecule(p)
    frame = _read_numeric_text(p)
    if frame.shape[1] >= 3 and str(frame.columns[0]) in ('0', 'col0'):
        frame.columns = (['longitude', 'latitude', 'depth']
                         + [f'col{i}' for i in range(frame.shape[1])])[:frame.shape[1]]
    return frame


# ------------------------------------------------------------ dive computers
def _dive_number(value):
    """Dive logs write '3.0 m', '25 C' and '0:10 min'; return the magnitude."""
    import re
    text = str(value).strip()
    clock = re.match(r'^(\d+):(\d{1,2})(?::(\d{1,2}))?', text)
    if clock:
        seconds = 0
        for part in [int(g) for g in clock.groups() if g is not None]:
            seconds = seconds * 60 + part
        return float(seconds)
    number = re.match(r'^[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?', text)
    return float(number.group(0)) if number else value


def _dive_frame(rows, what):
    if not rows:
        raise ImportError_(f"No {what} were found in this dive log.")
    frame = pd.DataFrame(rows)
    for col in frame.columns:
        if col != 'dive':
            frame[col] = frame[col].map(_dive_number)
    return frame


def _r_uddf(p):
    """Universal Dive Data Format - the interchange format Subsurface, MacDive,
    DiveLogs.de and most dive computers can export."""
    import gzip
    import xml.etree.ElementTree as ET
    with open(str(p), 'rb') as fh:
        raw = fh.read()
    if raw[:2] == b'\x1f\x8b':
        raw = gzip.decompress(raw)
    root = ET.fromstring(raw)
    rows, dive = [], 0
    for el in root.iter():
        tag = el.tag.split('}')[-1]
        if tag == 'dive':
            dive += 1
        elif tag == 'waypoint':
            row = {'dive': dive}
            for child in el:
                text = (child.text or '').strip()
                if text:
                    row[child.tag.split('}')[-1]] = text
            if len(row) > 1:
                rows.append(row)
    return _dive_frame(rows, 'profile waypoints')


def _r_ssrf(p):
    """Subsurface dive log (.ssrf): one <sample> per profile point."""
    import xml.etree.ElementTree as ET
    root = ET.parse(str(p)).getroot()
    rows, dive = [], 0
    for el in root.iter():
        tag = el.tag.split('}')[-1]
        if tag == 'dive':
            dive += 1
        elif tag == 'sample':
            row = {'dive': dive}
            row.update(el.attrib)
            rows.append(row)
    return _dive_frame(rows, 'profile samples')


def _r_dl7(p):
    """DAN DL7 dive logs (.zxu, .zxl, .dl7): samples sit in ZDP blocks."""
    rows, inside, dive = [], False, 0
    with open(str(p), 'r', errors='replace') as fh:
        for line in fh:
            s = line.strip()
            if s.startswith('ZDP{'):
                inside, dive = True, dive + 1
                continue
            if s.startswith('ZDP}'):
                inside = False
                continue
            if inside and s:
                parts = s.split('|')
                row = {'dive': dive}
                for i, name in enumerate(('time_minutes', 'depth', 'gas_pressure',
                                          'temperature', 'ascent_rate')):
                    if i < len(parts) and parts[i].strip():
                        row[name] = parts[i].strip()
                if len(row) > 1:
                    rows.append(row)
    return _dive_frame(rows, 'ZDP profile records')


def _r_sml(p):
    """Suunto Movescount / DM5 export (.sml): one <Sample> per profile point."""
    import xml.etree.ElementTree as ET
    root = ET.parse(str(p)).getroot()
    rows = []
    for el in root.iter():
        if el.tag.split('}')[-1] != 'Sample':
            continue
        row = {}
        for child in el:
            text = (child.text or '').strip()
            if text:
                row[child.tag.split('}')[-1]] = text
        if row:
            rows.append(row)
    return _dive_frame(rows, 'samples')


def _r_garmin_fit(p):
    """Garmin and Shearwater FIT activity and dive logs."""
    _need('fitparse', 'fitparse')
    from fitparse import FitFile
    fit = FitFile(str(p))
    rows = [{d.name: d.value for d in rec} for rec in fit.get_messages('record')]
    if not rows:
        rows = [{d.name: d.value for d in rec} for rec in fit.get_messages()]
    if not rows:
        raise ImportError_("No records were found in this FIT file.")
    return pd.DataFrame(rows)


def _r_fit_dispatch(p):
    """.fit is claimed by both FITS astronomy images and Garmin/Shearwater
    activity and dive logs, so decide on content: a FIT file carries the ASCII
    tag '.FIT' at byte 8, a FITS file starts 'SIMPLE  ='."""
    with open(str(p), 'rb') as fh:
        head = fh.read(12)
    if head[8:12] == b'.FIT':
        return _r_garmin_fit(p)
    return _r_fits(p)


# ------------------------------------------------------ genomics and sequences
def _r_vcf(p):
    """Variant Call Format: '##' metadata then a '#CHROM' header row."""
    import io as _io
    kept, header = [], None
    with open(str(p), 'r', errors='replace') as fh:
        for line in fh:
            if line.startswith('##'):
                continue
            if line.startswith('#CHROM'):
                header = line[1:]
                continue
            kept.append(line)
    if header is None:
        raise ImportError_("This VCF has no #CHROM header line.")
    return pd.read_csv(_io.StringIO(header + ''.join(kept)), sep='\t')


_INTERVAL_COLUMNS = {
    '.bed': ['chrom', 'start', 'end', 'name', 'score', 'strand'],
    '.gff': ['seqid', 'source', 'type', 'start', 'end', 'score',
             'strand', 'phase', 'attributes'],
    '.sam': ['qname', 'flag', 'rname', 'pos', 'mapq', 'cigar', 'rnext',
             'pnext', 'tlen', 'seq', 'qual'],
}


def _r_intervals(p):
    """BED, GFF/GFF3, GTF and SAM - tab tables with fixed, unwritten columns."""
    import io as _io
    ext = os.path.splitext(str(p).lower())[1]
    if ext in ('.gff3', '.gtf'):
        ext = '.gff'
    names = _INTERVAL_COLUMNS.get(ext, _INTERVAL_COLUMNS['.gff'])

    # The header lines have to go before read_csv sees the file, not after.
    # A SAM header is '@HD\tVN:1.6\tSO:coordinate' - three fields - and
    # read_csv infers its column count from the first line it reads, so with
    # on_bad_lines='skip' every real eleven-field alignment row was discarded
    # as malformed and the import failed with "no numeric columns to plot".
    # BED's optional 'track' and 'browser' lines have the same effect.
    body = []
    with open(str(p), 'r', errors='replace') as fh:
        for line in fh:
            if line.startswith(('#', '@', 'track', 'browser')):
                continue
            if not line.strip():
                continue
            body.append(line)
    if not body:
        raise ImportError_("No data lines were found in this file.")

    frame = pd.read_csv(_io.StringIO(''.join(body)), sep='\t', header=None,
                        names=names, engine='python', on_bad_lines='skip')
    frame = frame.dropna(axis=1, how='all')
    return frame


def _r_fasta(p):
    """FASTA: one row per record with its length and base composition."""
    rows, name, seq = [], None, []

    def flush():
        if name is None:
            return
        text = ''.join(seq).upper()
        row = {'id': name.split()[0] if name.split() else name,
               'description': name, 'length': len(text)}
        for base in 'ACGTUN':
            row[f'count_{base}'] = text.count(base)
        row['gc_fraction'] = ((text.count('G') + text.count('C')) / len(text)
                              if text else np.nan)
        rows.append(row)

    with open(str(p), 'r', errors='replace') as fh:
        for line in fh:
            if line.startswith('>'):
                flush()
                name, seq = line[1:].strip(), []
            elif name is not None:
                seq.append(line.strip())
    flush()
    if not rows:
        raise ImportError_("No FASTA records were found.")
    return pd.DataFrame(rows)


def _r_fastq(p):
    """FASTQ: one row per read with its length and mean Phred quality."""
    rows = []
    with open(str(p), 'r', errors='replace') as fh:
        while True:
            head = fh.readline()
            if not head:
                break
            seq = fh.readline().strip()
            fh.readline()
            qual = fh.readline().strip()
            if not qual:
                break
            scores = [ord(c) - 33 for c in qual]
            upper = seq.upper()
            rows.append({'id': head[1:].split()[0] if head.startswith('@') else head.strip(),
                         'length': len(seq),
                         'mean_quality': float(np.mean(scores)) if scores else np.nan,
                         'min_quality': int(min(scores)) if scores else 0,
                         'gc_fraction': ((upper.count('G') + upper.count('C')) / len(upper)
                                         if upper else np.nan)})
    if not rows:
        raise ImportError_("No FASTQ reads were found.")
    return pd.DataFrame(rows)


# -------------------------------------------------------- molecular structures
def _r_xyz_molecule(p):
    """Molecular XYZ, including multi-frame trajectories."""
    rows, frame_index, i = [], 0, 0
    with open(str(p), 'r', errors='replace') as fh:
        lines = fh.read().splitlines()
    while i < len(lines):
        head = lines[i].strip()
        if not head.isdigit():
            i += 1
            continue
        count = int(head)
        i += 2
        for j in range(count):
            if i + j >= len(lines):
                break
            parts = lines[i + j].split()
            if len(parts) >= 4:
                try:
                    rows.append({'frame': frame_index, 'element': parts[0],
                                 'x': float(parts[1]), 'y': float(parts[2]),
                                 'z': float(parts[3])})
                except ValueError:
                    pass
        i += count
        frame_index += 1
    if not rows:
        raise ImportError_("No atom records were found in this XYZ file.")
    return pd.DataFrame(rows)


def _r_pdb(p):
    """Protein Data Bank coordinates: the fixed-column ATOM/HETATM records."""
    rows = []
    with open(str(p), 'r', errors='replace') as fh:
        for line in fh:
            if not line.startswith(('ATOM', 'HETATM')):
                continue
            try:
                rows.append({
                    'record': line[0:6].strip(), 'serial': int(line[6:11]),
                    'atom': line[12:16].strip(), 'residue': line[17:20].strip(),
                    'chain': line[21:22].strip(), 'residue_number': int(line[22:26]),
                    'x': float(line[30:38]), 'y': float(line[38:46]),
                    'z': float(line[46:54]),
                    'occupancy': float(line[54:60] or 'nan'),
                    'b_factor': float(line[60:66] or 'nan'),
                    'element': line[76:78].strip(),
                })
            except ValueError:
                continue
    if not rows:
        raise ImportError_("No ATOM or HETATM records were found.")
    return pd.DataFrame(rows)


def _r_cif(p):
    """mmCIF/CIF: read the _atom_site loop_ block."""
    with open(str(p), 'r', errors='replace') as fh:
        lines = fh.read().splitlines()
    columns, rows = [], []
    for line in lines:
        s = line.strip()
        if s.startswith('loop_'):
            if rows:
                break
            columns = []
            continue
        if s.startswith('_atom_site.'):
            columns.append(s.split('.', 1)[1].split()[0])
            continue
        if columns and s and not s.startswith(('_', '#')):
            # mmCIF quotes any value containing a space, so a plain .split()
            # gives the wrong field count for rows like  'CA A'  - and the
            # `break` below then abandoned the whole loop at the first one,
            # returning the atoms before it as if that were the structure.
            try:
                parts = shlex.split(s)
            except ValueError:
                parts = s.split()
            if len(parts) == len(columns):
                rows.append(parts)
            continue
        if rows:
            break
    if not rows:
        raise ImportError_("No _atom_site loop was found in this CIF file.")
    return pd.DataFrame(rows, columns=columns)


def _r_molfile(p):
    """MDL molfile / SDF: the counts line, then the atom block."""
    rows, record, i = [], 0, 0
    with open(str(p), 'r', errors='replace') as fh:
        lines = fh.read().splitlines()
    while i + 3 < len(lines):
        try:
            atoms = int(lines[i + 3][0:3])
        except (ValueError, IndexError):
            i += 1
            continue
        record += 1
        for j in range(atoms):
            k = i + 4 + j
            if k >= len(lines):
                break
            parts = lines[k].split()
            if len(parts) >= 4:
                try:
                    rows.append({'record': record, 'x': float(parts[0]),
                                 'y': float(parts[1]), 'z': float(parts[2]),
                                 'element': parts[3]})
                except ValueError:
                    pass
        end = next((n for n in range(i, len(lines)) if lines[n].startswith('$$$$')), None)
        if end is None:
            break
        i = end + 1
    if not rows:
        raise ImportError_("No atom block was found in this molfile.")
    return pd.DataFrame(rows)


# ------------------------------------------------------------------- registry
READERS: dict[str, tuple] = {}
def _register(exts, fn, group):
    for e in exts:
        READERS[e] = (fn, group)

_register(('.csv',), _r_csv, 'Tables')
_register(('.tsv',), _r_tsv, 'Tables')
_register(('.txt', '.asc', '.dat', '.prn', '.log'), _r_sniff, 'Tables')
_register(('.json',), _r_json, 'Tables')
_register(('.jsonl', '.ndjson'), _r_jsonl, 'Tables')
_register(('.xlsx', '.xls', '.xlsm'), _r_excel, 'Spreadsheets')
_register(('.xlsb',), _r_xlsb, 'Spreadsheets')
_register(('.ods',), _r_ods, 'Spreadsheets')
_register(('.html', '.htm'), _r_html, 'Tables')
_register(('.xml',), _r_xml, 'Tables')
_register(('.yaml', '.yml'), _r_yaml, 'Tables')
_register(('.toml',), _r_toml, 'Tables')
_register(('.parquet', '.pq'), _r_parquet, 'Columnar')
_register(('.feather',), _r_feather, 'Columnar')
_register(('.orc',), _r_orc, 'Columnar')
_register(('.arrow', '.ipc'), _r_arrow, 'Columnar')
_register(('.sav', '.por'), _r_spss, 'Statistics')
_register(('.dta',), _r_stata, 'Statistics')
_register(('.sas7bdat', '.xpt'), _r_sas, 'Statistics')
_register(('.npy',), _r_npy, 'Arrays')
_register(('.npz',), _r_npz, 'Arrays')
_register(('.pkl', '.pickle'), _r_pickle, 'Arrays')
_register(('.db', '.sqlite', '.sqlite3'), _r_sqlite, 'Databases')
_register(('.mat',), _r_matlab, 'MATLAB')
_register(('.h5', '.hdf', '.hdf5', '.he5'), _r_hdf, 'HDF5')
_register(('.nc', '.netcdf', '.nc4', '.cdf'), _r_netcdf, 'NetCDF')
_register(('.tdms',), _r_tdms, 'Instruments')
_register(('.tdm',), _r_tdm, 'Instruments')
_register(('.wav', '.flac', '.ogg', '.aiff', '.aif'), _r_audio, 'Audio')
_register(('.las',), _r_las, 'Well logs')
_register(('.mgf', '.mzml', '.mzxml'), _r_massspec, 'Mass spectrometry')
_register(('.fcs',), _r_fcs, 'Flow cytometry')
_register(('.jdx', '.dx', '.jcamp'), _r_jcamp, 'Spectroscopy')
_register(('.fits', '.fit', '.fts'), _r_fits, 'Astronomy')
_register(('.root',), _r_root, 'Particle physics')
_register(('.mseed', '.miniseed', '.sac', '.segy', '.sgy'), _r_seismic, 'Seismic')
_register(('.edf',), _r_edf, 'Physiology')
_register(('.bdf', '.fif', '.set', '.cnt'), _r_mne, 'Physiology')
_register(('.abf',), _r_abf, 'Electrophysiology')
_register(('.nwb',), _r_nwb, 'Neuroscience')
_register(('.dcm', '.dicom'), _r_dicom, 'Medical imaging')
_register(('.nii',), _r_nifti, 'Medical imaging')
_register(('.tif', '.tiff'), _r_tiff, 'Imaging')
_register(('.geojson', '.shp', '.gpkg'), _r_geo, 'Geospatial')

# Business, finance and reporting
_register(('.dbf',), _r_dbf, 'Business')
_register(('.mdb', '.accdb'), _r_access, 'Business')
_register(('.numbers',), _r_numbers, 'Spreadsheets')
_register(('.duckdb', '.ddb'), _r_duckdb, 'Databases')
_register(('.avro',), _r_avro, 'Columnar')
_register(('.msgpack', '.mpk'), _r_msgpack, 'Arrays')
_register(('.rds', '.rdata', '.rda'), _r_rdata, 'Statistics')
_register(('.fwf',), _r_fixed_width, 'Tables')
_register(('.psv',), _r_psv, 'Tables')
_register(('.ofx', '.qfx'), _r_ofx, 'Finance')
_register(('.qif',), _r_qif, 'Finance')
_register(('.sta', '.mt940'), _r_mt940, 'Finance')
_register(('.pdf',), _r_pdf_tables, 'Reports')
_register(('.docx',), _r_docx_tables, 'Reports')

# Marine, oceanography, hydrography and diving
_register(('.cnv', '.btl', '.ctd'), _r_seabird, 'Marine')
_register(('.odv',), _r_odv, 'Marine')
_register(('.rsk',), _r_rsk, 'Marine')
_register(('.pd0', '.000', '.vec', '.aqd', '.wpr', '.prf'), _r_adcp, 'Marine')
_register(('.nmea', '.gps'), _r_nmea, 'Marine')
_register(('.grib', '.grib2', '.grb'), _r_grib, 'Marine')
_register(('.xyz',), _r_xyz, 'Marine')
_register(('.gpx',), _r_gpx, 'Tracks')
_register(('.kml', '.kmz'), _r_kml, 'Tracks')
_register(('.uddf',), _r_uddf, 'Diving')
_register(('.ssrf',), _r_ssrf, 'Diving')
_register(('.zxu', '.zxl', '.dl7'), _r_dl7, 'Diving')
_register(('.sml',), _r_sml, 'Diving')
# .fit is claimed by FITS and by Garmin/Shearwater dive logs; _r_fit_dispatch
# decides on the file's own magic bytes, so it stays listed under Astronomy.
_register(('.fit',), _r_fit_dispatch, 'Astronomy')

# Genomics and molecular structures
_register(('.vcf',), _r_vcf, 'Genomics')
_register(('.bed', '.gff', '.gff3', '.gtf', '.sam'), _r_intervals, 'Genomics')
_register(('.fasta', '.fa', '.fna'), _r_fasta, 'Genomics')
_register(('.fastq', '.fq'), _r_fastq, 'Genomics')
_register(('.pdb',), _r_pdb, 'Molecular structures')
_register(('.cif',), _r_cif, 'Molecular structures')
_register(('.mol', '.sdf'), _r_molfile, 'Molecular structures')

ALL_EXT = tuple(sorted(READERS))


def format_groups() -> dict[str, list[str]]:
    groups: dict[str, list[str]] = {}
    for ext, (_fn, group) in READERS.items():
        groups.setdefault(group, []).append(ext)
    return {g: sorted(v) for g, v in sorted(groups.items())}


_COMPRESSED = ('.gz', '.bz2', '.xz', '.zst')


def _decompressed_copy(src: Path):
    """Return (path_to_read, temp_path_or_None) for a possibly compressed file.

    _suffix() strips the compression suffix so 'sample.vcf.gz' correctly routes
    to the VCF reader - but only the pandas-backed readers decompress for
    themselves. About thirty readers here open the file with a plain open(), so
    a gzipped VCF, FASTA, FASTQ, GPX or Sea-Bird cast reached them as a raw
    deflate stream and failed with a message about the format ("This VCF has no
    #CHROM header line") that pointed nowhere near the real cause.

    Decompressing to a temp file for every reader is a few hundred milliseconds
    on a large file and removes the whole class of problem; the pandas readers
    are perfectly happy with an already-decompressed file.
    """
    low = str(src).lower()
    suffix = next((z for z in _COMPRESSED if low.endswith(z)), None)
    if suffix is None:
        return str(src), None

    if suffix == '.gz':
        import gzip
        opener = lambda: gzip.open(str(src), 'rb')
    elif suffix == '.bz2':
        import bz2
        opener = lambda: bz2.open(str(src), 'rb')
    elif suffix == '.xz':
        import lzma
        opener = lambda: lzma.open(str(src), 'rb')
    else:
        try:
            import zstandard  # noqa: F401
        except ImportError:
            raise ImportError_(
                "Reading .zst files needs the 'zstandard' package. "
                "Run INSTALL-DATA-FORMATS.bat, or decompress the file first."
            )
        import zstandard as zstd
        def opener():
            return zstd.ZstdDecompressor().stream_reader(open(str(src), 'rb'))

    inner = os.path.splitext(low[: -len(suffix)])[1] or '.dat'
    fd, temp = tempfile.mkstemp(suffix=inner, prefix='graphvis-')
    os.close(fd)
    try:
        with opener() as source, open(temp, 'wb') as sink:
            shutil.copyfileobj(source, sink, length=1 << 20)
    except ImportError_:
        raise
    except Exception as exc:
        try:
            os.remove(temp)
        except OSError:
            pass
        raise ImportError_(f"Could not decompress this {suffix} file: {exc}")
    return temp, temp


def _suffix(path: str) -> str:
    low = path.lower()
    # pandas transparently decompresses these, so the real format is the one
    # before the compression suffix.
    for zsuf in ('.gz', '.bz2', '.xz', '.zip', '.zst'):
        if low.endswith(zsuf):
            low = low[: -len(zsuf)]
            break
    if low.endswith('.nii'):
        return '.nii'
    return os.path.splitext(low)[1]


def import_to_arrow(path: str, out_dir: str) -> dict:
    """Convert any supported dataset into Arrow IPC. Returns a summary dict."""
    src = Path(path)
    if not src.exists():
        raise ImportError_(f"File not found: {path}")

    ext = _suffix(str(src))
    entry = READERS.get(ext)
    if entry is None:
        raise ImportError_(
            f"Unsupported dataset format '{ext or src.name}'. "
            f"{len(ALL_EXT)} formats are supported; see the Add dataset dialog."
        )
    reader, group = entry
    read_path, temp_path = _decompressed_copy(src)
    try:
        frame = reader(read_path)
    finally:
        if temp_path:
            try:
                os.remove(temp_path)
            except OSError:
                pass

    if frame is None or len(frame) == 0:
        raise ImportError_("No usable data was found in this file.")

    frame = pd.DataFrame(frame)
    frame.columns = [str(c) for c in frame.columns]
    for col in frame.columns:
        # pandas 3 gives text columns a StringDtype rather than object, so the
        # old `dtype == object` test silently stopped matching and every reader
        # that returns text failed later with "no numeric columns". Ask the
        # question we actually mean: is this column not already numeric?
        if pd.api.types.is_numeric_dtype(frame[col]):
            continue
        if pd.api.types.is_datetime64_any_dtype(frame[col]):
            continue
        try:
            converted = pd.to_numeric(frame[col], errors='coerce')
        except (TypeError, ValueError):
            continue
        if converted.notna().mean() > 0.8:
            frame[col] = converted

    numeric = [c for c in frame.columns if pd.api.types.is_numeric_dtype(frame[c])]
    if not numeric:
        raise ImportError_("This file has no numeric columns to plot.")

    import pyarrow as pa
    import pyarrow.ipc as ipc

    os.makedirs(out_dir, exist_ok=True)
    stem = src.name
    for suf in ('.gz', '.bz2', '.xz', '.zip', '.zst'):
        if stem.lower().endswith(suf):
            stem = stem[: -len(suf)]
    stem = os.path.splitext(stem)[0]
    out_path = os.path.join(out_dir, f"{stem}.arrow")

    table = pa.Table.from_pandas(frame, preserve_index=False)
    with pa.OSFile(out_path, 'wb') as sink:
        with ipc.new_file(sink, table.schema) as writer:
            writer.write_table(table)

    return {
        'ok': True,
        'arrow_path': out_path,
        'name': stem,
        'kind': group,
        'rows': int(len(frame)),
        'columns': list(frame.columns),
        'numeric_columns': numeric,
    }
