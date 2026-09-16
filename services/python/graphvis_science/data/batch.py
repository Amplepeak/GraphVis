"""Folder-scale batch processing and report generation.

A port of GraphVis 17's ``automation/batch.py``. Point it at a folder, run one
operation over every dataset in it, and get an HTML, Word or PDF report of what
succeeded and what did not.

Two things changed in the port.

The file filter is no longer a hard-coded set of 26 extensions. v18's importer
knows 149 formats and publishes them, so the scan asks the importer what it can
read. v17's list had drifted out of date the moment a reader was added.

The operation is named rather than passed in. v17 took a Python callable, which
a native caller cannot supply across a JSON pipe. The three that matter -
summarise, describe, convert - are here by name, and each one reports per file
rather than aborting the run.
"""
from __future__ import annotations

import html
import json
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

OPERATIONS = ("summary", "describe", "convert")


@dataclass(slots=True)
class BatchItemResult:
    path: str
    ok: bool
    summary: str
    payload: Any = None


@dataclass(slots=True)
class BatchRunResult:
    items: list[BatchItemResult] = field(default_factory=list)
    output_files: list[str] = field(default_factory=list)

    @property
    def successes(self) -> int:
        return sum(1 for i in self.items if i.ok)

    @property
    def failures(self) -> int:
        return len(self.items) - self.successes


def readable_extensions() -> set[str]:
    """Every extension the importer can read, asked rather than assumed."""
    from graphvis_science.data.importer import ALL_EXT
    return {e if e.startswith(".") else "." + e for e in ALL_EXT}


def scan(folder: str, recursive: bool = False) -> list[str]:
    root = Path(folder)
    if not root.is_dir():
        raise NotADirectoryError(f"Not a folder: {folder}")
    known = readable_extensions()
    walker = root.rglob("*") if recursive else root.glob("*")
    found = []
    for p in walker:
        if not p.is_file():
            continue
        name = p.name.lower()
        # A compressed dataset keeps its real extension in front of the
        # compression suffix, and the importer decompresses on the way in.
        for squeeze in (".gz", ".bz2", ".xz", ".zst"):
            if name.endswith(squeeze):
                name = name[: -len(squeeze)]
                break
        if os.path.splitext(name)[1] in known:
            found.append(str(p))
    return sorted(found)


def run(paths: Iterable[str], operation: str = "summary", out_dir: str = "") -> BatchRunResult:
    from graphvis_science.data.importer import import_to_arrow

    if operation not in OPERATIONS:
        raise ValueError(f"Unknown batch operation '{operation}'. "
                         f"Choose one of: {', '.join(OPERATIONS)}.")
    staging = out_dir or os.path.join(os.path.dirname(str(next(iter(paths), "."))), "_batch")
    result = BatchRunResult()

    for path in paths:
        try:
            info = import_to_arrow(path, staging)
            rows = int(info.get("rows", 0))
            columns = list(info.get("columns", []))
            summary = f"{rows} rows x {len(columns)} columns ({info.get('kind', 'data')})"

            if operation == "summary":
                payload = {"rows": rows, "columns": columns, "kind": info.get("kind", "")}
            elif operation == "convert":
                payload = {"arrow_path": info.get("arrow_path", ""), "rows": rows}
            else:
                import pyarrow as pa
                import pyarrow.ipc as ipc
                with pa.memory_map(str(info["arrow_path"]), "r") as source:
                    frame = ipc.open_file(source).read_all().to_pandas()
                numeric = frame.select_dtypes("number")
                payload = json.loads(numeric.describe().to_json()) if not numeric.empty else {}

            result.items.append(BatchItemResult(path, True, summary, payload))
        except Exception as exc:  # one bad file must not end the run
            result.items.append(BatchItemResult(path, False, f"{type(exc).__name__}: {exc}"))

    return result


def write_html_report(result: BatchRunResult, path: str,
                      title: str = "GraphVis Batch Report") -> str:
    rows = []
    for item in result.items:
        payload = (html.escape(json.dumps(item.payload, default=str, ensure_ascii=False)[:4000])
                   if item.ok else "")
        rows.append(
            f"<tr><td>{html.escape(Path(item.path).name)}</td>"
            f"<td class='{'ok' if item.ok else 'bad'}'>{'OK' if item.ok else 'ERROR'}</td>"
            f"<td>{html.escape(item.summary)}</td><td><pre>{payload}</pre></td></tr>"
        )
    doc = (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<title>{html.escape(title)}</title><style>"
        "body{font-family:Arial,Helvetica,sans-serif;margin:2em;color:#222}"
        "table{border-collapse:collapse;width:100%}"
        "td,th{border:1px solid #ccc;padding:.5em;vertical-align:top;text-align:left}"
        "th{background:#f2f2f2}pre{white-space:pre-wrap;margin:0;font-size:.85em}"
        ".ok{color:#1a7f37;font-weight:bold}.bad{color:#b3261e;font-weight:bold}"
        "</style></head><body>"
        f"<h1>{html.escape(title)}</h1>"
        f"<p>{result.successes} succeeded, {result.failures} failed.</p>"
        "<table><tr><th>File</th><th>Status</th><th>Summary</th><th>Payload</th></tr>"
        f"{''.join(rows)}</table></body></html>"
    )
    Path(path).write_text(doc, encoding="utf-8")
    result.output_files.append(path)
    return path


def write_word_report(result: BatchRunResult, path: str,
                      title: str = "GraphVis Batch Report") -> str:
    try:
        from docx import Document
    except Exception as exc:
        raise RuntimeError("Word reports need python-docx, which is part of the "
                           "science add-on's [reports] group.") from exc
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"{result.successes} succeeded, {result.failures} failed.")
    table = doc.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text = "File", "Status", "Summary"
    for item in result.items:
        cells = table.add_row().cells
        cells[0].text = Path(item.path).name
        cells[1].text = "OK" if item.ok else "ERROR"
        cells[2].text = item.summary
    doc.save(path)
    result.output_files.append(path)
    return path


def write_pdf_report(result: BatchRunResult, path: str,
                     title: str = "GraphVis Batch Report") -> str:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                        Table, TableStyle)
    except Exception as exc:
        raise RuntimeError("PDF reports need reportlab, which is part of the "
                           "science add-on's [reports] group.") from exc
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12),
             Paragraph(f"{result.successes} succeeded, {result.failures} failed.",
                       styles["BodyText"]),
             Spacer(1, 12)]
    data = [["File", "Status", "Summary"]]
    data += [[Path(i.path).name, "OK" if i.ok else "ERROR", i.summary[:180]]
             for i in result.items]
    table = Table(data, repeatRows=1, colWidths=[140, 55, 300])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    SimpleDocTemplate(path, pagesize=A4).build(story)
    result.output_files.append(path)
    return path


def write_report(result: BatchRunResult, path: str, fmt: str = "html",
                 title: str = "GraphVis Batch Report") -> str:
    writers = {"html": write_html_report, "docx": write_word_report,
               "word": write_word_report, "pdf": write_pdf_report}
    writer = writers.get(str(fmt).lower())
    if writer is None:
        raise ValueError(f"Unknown report format '{fmt}'. Choose html, docx or pdf.")
    return writer(result, path, title)
